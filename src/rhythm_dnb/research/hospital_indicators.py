"""Audit hospital indicator coverage and participant-resampled networks without inventing outcomes."""

import csv
import json
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path

import numpy as np

from ..io.sources.hospital_workbook import CLOCKS, FIELDS, META, normalized_number, read_monitoring_workbook
from ..provenance import file_hash


def validate_plan(plan):
    # PSEUDOCODE: reject unsupported statistics or panels rather than silently changing the analysis.
    if plan.get('purpose') != 'descriptive_hospital_indicator_audit_not_warning_validation':
        raise ValueError('This analysis cannot validate warning predictions.')
    if set(plan.get('panels', {})) != {'core', 'expanded'}:
        raise ValueError('Declare core and expanded panels.')
    eligible = set(FIELDS) - set(META) - set(CLOCKS) - {'exercise_type', 'nap_county', 'late_night_eating_flag', 'smoking_state', 'self_rate_state', 'time_in_bed'}
    for features in plan['panels'].values():
        if len(features) < 2 or len(features) != len(set(features)) or not set(features) <= eligible:
            raise ValueError('Panels require unique eligible numeric measurements.')
    if not set(plan['panels']['core']) < set(plan['panels']['expanded']):
        raise ValueError('Core must be a strict subset of expanded.')
    if type(plan.get('bootstrap_replicates')) is not int or plan['bootstrap_replicates'] < 100:
        raise ValueError('At least 100 bootstrap draws are required for this descriptive interval output.')
    if type(plan.get('seed')) is not int or plan['seed'] < 0 or plan.get('interval_quantiles') != [.025, .975]:
        raise ValueError('Specify a nonnegative seed and the supported pointwise 95% interval.')


def participant_cube(rows, features, *, include_questioned=False):
    """Use identical complete people at every occasion, retaining exclusion reasons."""
    # PSEUDOCODE: group complete person blocks -> align source occasions -> emit person-by-occasion-by-node array.
    grouped = defaultdict(dict)
    for row in rows:
        pid, occasion = row['participant_id'], row['occasion']
        if occasion in grouped[pid]:
            raise ValueError('Duplicate person/occasion; never silently collapse visits.')
        grouped[pid][occasion] = row
    occasions = sorted({r['occasion'] for r in rows})
    people, blocks, excluded = [], [], []
    for pid, visits in sorted(grouped.items()):
        block, reasons = [], []
        for occasion in occasions:
            if occasion not in visits:
                reasons.append(f'occasion_{occasion}:not_collected')
                continue
            row, values = visits[occasion], []
            for field in features:
                value = row['values'].get(field)
                if include_questioned and row['states'].get(field) == 'requires_confirmation':
                    value = normalized_number(field, row['raw'][field]['value'])
                if value is None or not np.isfinite(value):
                    reasons.append(f'occasion_{occasion}:{field}:{row["states"].get(field, "not_collected")}')
                values.append(value)
            block.append(values)
        if reasons:
            excluded.append({'participant_id': pid, 'reasons': reasons})
        else:
            people.append(pid)
            blocks.append(block)
    return np.asarray(blocks, dtype=float).reshape(len(people), len(occasions), len(features)), people, occasions, excluded


def correlation_summary(matrix):
    """Participation ratio describes covariance concentration, not independently measured dimensions."""
    # PSEUDOCODE: check finite observations and variable columns -> compute Pearson matrix and its concentration.
    x = np.asarray(matrix, dtype=float)
    if x.ndim != 2 or x.shape[0] < 3 or x.shape[1] < 2 or not np.isfinite(x).all():
        return None
    sd = x.std(axis=0, ddof=1)
    if np.any(sd == 0):
        return None
    correlation = np.corrcoef(x, rowvar=False)
    if not np.isfinite(correlation).all():
        return None
    p = x.shape[1]
    eigenvalues = np.maximum(np.linalg.eigvalsh(correlation), 0)
    return {'correlation': correlation, 'sd': sd, 'participation_ratio': float(p * p / np.square(correlation).sum()),
            'largest_eigenvalue_share': float(eigenvalues[-1] / p)}


def summarize_panel(cube, features, occasions, draws=None):
    # PSEUDOCODE: retain every occasion and edge -> resample people -> return pointwise intervals and invalid counts.
    summaries, edges, matrices = [], [], {}
    upper = np.triu_indices(len(features), 1)
    for t, occasion in enumerate(occasions):
        point = correlation_summary(cube[:, t, :])
        if point is None:
            summaries.append({'occasion': occasion, 'people': len(cube), 'status': 'insufficient_or_constant_values'})
            continue
        r = point['correlation']; matrices[str(occasion)] = r.tolist()
        item = {'occasion': occasion, 'people': len(cube), 'status': 'descriptive_only',
                'features': len(features), 'edge_count': len(upper[0]),
                'participation_ratio': point['participation_ratio'],
                'largest_eigenvalue_share': point['largest_eigenvalue_share'],
                'mean_abs_correlation': float(np.abs(r[upper]).mean()),
                'sd_by_feature': dict(zip(features, point['sd'].tolist()))}
        sampled, ranks = [], []
        if draws is not None:
            for draw in draws:
                sample = correlation_summary(cube[draw, t, :])
                if sample is not None:
                    sampled.append(sample['correlation'][upper]); ranks.append(sample['participation_ratio'])
            item['bootstrap_valid_draws'] = len(sampled)
            item['bootstrap_total_draws'] = len(draws)
        # A failed draw is visible; do not present conditional-on-success intervals as unconditional intervals.
        complete = draws is not None and len(sampled) == len(draws)
        low, high = np.quantile(sampled, [.025, .975], axis=0) if complete else (None, None)
        item['participation_ratio_interval'] = np.quantile(ranks, [.025, .975]).tolist() if complete else None
        item['median_edge_interval_width'] = float(np.median(high-low)) if complete else None
        for index, (a, b) in enumerate(zip(*upper)):
            edges.append({'occasion': occasion, 'feature_a': features[a], 'feature_b': features[b],
                          'pearson': float(r[a, b]), 'low': float(low[index]) if complete else None,
                          'high': float(high[index]) if complete else None,
                          'bootstrap_valid_draws': len(sampled) if draws is not None else None})
        summaries.append(item)
    return {'features': features, 'occasions': summaries, 'correlations': matrices, 'edges': edges}


def analyze_panels(data, plan):
    # PSEUDOCODE: hold people and bootstrap draws fixed across panels; report source-value sensitivity separately.
    validate_plan(plan)
    expanded = plan['panels']['expanded']
    cube, people, occasions, excluded = participant_cube(data['rows'], expanded)
    if len(people) < 3:
        return {'status': 'insufficient_complete_people', 'people': len(people), 'excluded': excluded,
                'warning_accuracy': None, 'panels': {}}
    draws = np.random.default_rng(plan['seed']).integers(0, len(people), (plan['bootstrap_replicates'], len(people)))
    panels = {}
    for name, features in plan['panels'].items():
        indices = [expanded.index(f) for f in features]
        panels[name] = summarize_panel(cube[:, :, indices], features, occasions, draws)
    alternate, alternate_people, alt_occasions, alt_excluded = participant_cube(data['rows'], expanded, include_questioned=True)
    return {'status': 'descriptive_complete_warning_not_evaluable', 'people': len(people),
            'participant_ids': people, 'occasions': occasions, 'excluded': excluded,
            'warning_accuracy': None, 'lead_time': None, 'panels': panels,
            'questioned_value_sensitivity': {
                'description': 'Use the questioned raw number as recorded. Cohort size may change; not a replacement or a pure within-cohort value effect.',
                'people': len(alternate_people), 'excluded': alt_excluded,
                'expanded': summarize_panel(alternate, expanded, alt_occasions)},
            'readiness': {'population_dnb': 'needs verified comparable event-relative stages',
                          'single_sample_dnb': 'needs independently confirmed stable reference people and outcomes',
                          'longitudinal_dnb': 'needs full dates and substantially denser repeated observations',
                          'event_labels_supplied': False, 'stable_reference_labels_supplied': False}}


def _json(path, value):
    # PSEUDOCODE: create immutable valid JSON; refuse to overwrite an earlier artifact.
    with Path(path).open('x', encoding='utf-8') as stream:
        json.dump(value, stream, ensure_ascii=False, indent=2, allow_nan=False)
        stream.write('\n')


def _csv(path, rows, fields):
    # PSEUDOCODE: write a stable review table, quoting content and preserving missing values as blanks.
    with Path(path).open('x', encoding='utf-8-sig', newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=fields)
        writer.writeheader(); writer.writerows(rows)


def export_workbook(data, path):
    # PSEUDOCODE: make one input sheet; source doubts live in cell comments instead of extra sheets or fabricated answers.
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    workbook = Workbook(); sheet = workbook.active; sheet.title = '监测输入'
    fields = [f for f in FIELDS if f not in META]
    metadata = ['record_id', 'participant_id', 'occasion', 'date_original', 'month_day', 'calendar_date', 'age', 'gender']
    sheet.append(metadata + fields)
    for column, field in enumerate(metadata + fields, 1):
        sheet.cell(1, column).font = Font(bold=True, color='FFFFFF')
        sheet.cell(1, column).fill = PatternFill('solid', fgColor='245B6B')
        definition = data['definitions'].get(field)
        if definition:
            sheet.cell(1, column).comment = Comment(f'{definition["label"]}；{definition["source_unit"]}', '字段说明')
        sheet.column_dimensions[get_column_letter(column)].width = 21
    for i, row in enumerate(data['rows'], 2):
        values = [row['record_id'], row['participant_id'], row['occasion'], str(row['raw']['date']['value']),
                  row['date']['month_day'], row['date']['calendar_date'], row['values']['age'], row['values']['gender']]
        for field in fields:
            value = row['values'][field]
            if field in CLOCKS and value is not None:
                seconds = round(value * 60)
                hour, remaining = divmod(seconds, 3600); minute, second = divmod(remaining, 60)
                value = f'{hour:02d}:{minute:02d}:{second:02d}'
            values.append(value)
        sheet.append(values)
        sheet.cell(i, 4).comment = Comment('原值；' + row['date']['status'] + '；候选月日：' + ', '.join(row['date']['candidates']), '源数据核对')
        for c, field in enumerate(fields, len(metadata) + 1):
            if row['states'][field] != 'observed':
                sheet.cell(i, c).comment = Comment(f'{row["states"][field]}；原值={row["raw"][field]["value"]}；{row["sheet"]}!{row["raw"][field]["cell"]}', '源数据核对')
    sheet.freeze_panes = 'D2'; sheet.auto_filter.ref = sheet.dimensions
    workbook.save(path); workbook.close()


def plot_networks(result, output):
    # PSEUDOCODE: show all occasions on the same color scale and dimension bars with resampling intervals.
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    if not result['panels']:
        return
    expanded = result['panels']['expanded']
    matrices = expanded['correlations']
    if matrices:
        fig, axes = plt.subplots(1, len(matrices), figsize=(5 * len(matrices), 5), squeeze=False, layout='constrained')
        for ax, (occasion, values) in zip(axes.flat, matrices.items()):
            im = ax.imshow(values, cmap='RdBu_r', vmin=-1, vmax=1)
            ax.set_title(f'Recording occasion {occasion}')
            ax.set_xticks(range(len(expanded['features'])), expanded['features'], rotation=90, fontsize=6)
            ax.set_yticks(range(len(expanded['features'])), expanded['features'], fontsize=6)
        fig.colorbar(im, ax=list(axes.flat), shrink=.75, label='Pearson correlation')
        fig.suptitle(f'{result["people"]} matched people; recording order is NOT disease stage')
        fig.savefig(output / 'networks.png', dpi=160); plt.close(fig)
    fig, ax = plt.subplots(figsize=(8, 5), layout='constrained')
    for offset, (name, panel) in zip((-.18, .18), result['panels'].items()):
        records = [r for r in panel['occasions'] if r['status'] == 'descriptive_only']
        x = np.asarray([r['occasion'] for r in records]) + offset
        y = [r['participation_ratio'] for r in records]
        ax.bar(x, y, .36, label=f'{name} ({len(panel["features"])} measured nodes)')
        for xi, record in zip(x, records):
            interval = record['participation_ratio_interval']
            if interval:
                ax.vlines(xi, *interval, color='black', linewidth=1)
    ax.set_xticks(result['occasions']); ax.set_xlabel('Recording occasion (not disease stage)')
    ax.set_ylabel('Correlation participation ratio')
    ax.set_title('Covariance diversity; not a warning accuracy or sufficiency test')
    ax.legend(); fig.savefig(output / 'dimensions.png', dpi=160); plt.close(fig)


def run_hospital_analysis(source, description, plan, output, *, plots=False):
    """Freeze this descriptive plan, preserve cell lineage, then compute only supported statistics."""
    # PSEUDOCODE: freeze sources and plan -> audit input -> run matched-person comparisons -> export verified artifacts.
    validate_plan(plan)
    source, description, output = Path(source), Path(description), Path(output)
    hashes = {'workbook': file_hash(source), 'description': file_hash(description)}
    output.mkdir(parents=True, exist_ok=False)
    _json(output / 'protocol.json', {'created_at': datetime.now(timezone.utc).isoformat(), 'plan': plan,
        'source_hashes': hashes, 'source_files': {'workbook': str(source.resolve()), 'description': str(description.resolve())},
        'code_hashes': {str(p): file_hash(p) for p in (Path(__file__).resolve(), Path(__file__).resolve().parents[1] / 'io/sources/hospital_workbook.py')},
        'prior_inspection': 'Source fields and errors inspected before this plan; this is an exploratory audit, not prospective preregistration.'})
    data = read_monitoring_workbook(source, plan.get('review_cells', []))
    result = analyze_panels(data, plan)
    _json(output / 'normalized.json', data)
    _json(output / 'analysis.json', result)
    _csv(output / 'source-review.csv', data['issues'], ['record_id', 'participant_id', 'sheet', 'cell', 'field', 'code', 'raw_value'])
    edges = [{'panel': name, **edge} for name, panel in result['panels'].items() for edge in panel['edges']]
    _csv(output / 'network-edges.csv', edges, ['panel', 'occasion', 'feature_a', 'feature_b', 'pearson', 'low', 'high', 'bootstrap_valid_draws'])
    export_workbook(data, output / 'hospital-inputs.xlsx')
    # Empty outcome fields are requests for hospital evidence, not negative event labels.
    followup_fields = ['participant_id', 'calendar_year', 'timezone', 'stable_from', 'stable_through',
                       'event_status', 'event_onset', 'event_confirmed_at', 'followup_through', 'outcome_definition', 'evidence_reference']
    _csv(output / 'followup-request.csv', [{'participant_id': pid} for pid in sorted({r['participant_id'] for r in data['rows']})], followup_fields)
    if plots:
        plot_networks(result, output)
    write_summary(data, result, plan, output)
    if hashes != {'workbook': file_hash(source), 'description': file_hash(description)}:
        raise ValueError('A source changed during analysis; outputs must not be treated as a fixed dataset.')
    _json(output / 'manifest.json', {'source_hashes': hashes, 'files': {p.name: file_hash(p) for p in sorted(output.iterdir()) if p.is_file()}})
    manifest = json.loads((output / 'manifest.json').read_text(encoding='utf-8'))
    if not all(file_hash(output / name) == digest for name, digest in manifest['files'].items()):
        raise ValueError('Output hash verification failed.')
    return {'output': str(output.resolve()), 'records': data['record_count'], 'people': data['participant_count'],
            'matched_people': result['people'], 'warning_accuracy': None, 'verified_files': len(manifest['files'])}


def write_summary(data, result, plan, output):
    # PSEUDOCODE: report accomplished checks and limited interpretation without presenting network structure as prediction accuracy.
    issue_counts = Counter(row['code'] for row in data['issues'])
    lines = ['# 医院指标检查结果', '',
        f'已整理 {data["participant_count"]} 人、{data["record_count"]} 条记录，合并为一个工作表。原始 XLS/DOCX 未修改；按各表列名读取，避免第 13 张表插列导致错位。', '',
        f'日期待核对：{issue_counts["ambiguous_day"]} 个数字日期存在日号歧义，{issue_counts["clock_in_date_cell"]} 个日期格被存为时刻；全部缺少明确年份。数值待核对：{issue_counts["invalid_value"]} 个越界或无效值，{issue_counts["requires_confirmation"]} 个疑似录入值。详情见 source-review.csv。没有自行补日期、结局或把未知填成零。', '',
        f'比较方案在计算前保存：{len(plan["panels"]["core"])} 项核心指标、{len(plan["panels"]["expanded"])} 项扩展指标，按同一批参与者比较。保留全部输入；未纳入初始 Pearson 网络的时刻、类别与等级字段仍在表内。', '',
        f'两套面板共同完整的参与者为 {result["people"]} 人。每次计算每人仅一条记录；按人整体重复抽样 {plan["bootstrap_replicates"]} 次，保留各次记录的关联。记录次序不是紊乱进展阶段。', '',
        '| 记录次序 | 核心指标相关矩阵的有效维数 | 扩展指标相关矩阵的有效维数 |', '|---|---:|---:|']
    if result['panels']:
        for a, b in zip(result['panels']['core']['occasions'], result['panels']['expanded']['occasions']):
            def shown(record):
                # PSEUDOCODE: leave unavailable estimates visibly unavailable.
                return f'{record["participation_ratio"]:.2f}' if record['status'] == 'descriptive_only' else '不可计算'
            lines.append(f'| {a["occasion"]} | {shown(a)} | {shown(b)} |')
    lines += ['', '这里的“有效维数”表示信息是否集中在少数共同变化方向，不等于独立生理指标数量，更不是正确率，也不能据此认定维度已经足够。扩展项有无预警增益仍需结局验证；各相关系数的区间是逐项描述区间，不能当成全网络显著性检验。', '',
        '可查看 dimensions.png 和 networks.png（启用 --plots 时生成）；完整数值在 analysis.json / network-edges.csv。原始可疑饮水量也单独保留了一次敏感性计算，未据此选择更好看的方案。', '',
        '下一份实测数据应补：明确年份与时区、稳定期依据、是否发生节律紊乱及发生/确认时间、随访截止时间和连续监测记录。followup-request.csv 已列好全部人员，未知字段留空，交由持有原始记录的人填写。', '',
        '**当前没有真实预警正确率。** 群体 DNB 还缺可比的事件前阶段，个体 sDNB 缺经确认的稳定参考人群，个人纵向 DNB 缺连续时间序列。四个离散测量日不能补成完整随访；医院数据也不自动等于健康参考数据。', '',
        '计算原则参考 [经典 DNB](https://www.nature.com/articles/srep00342) 与 [单样本 DNB](https://journals.plos.org/ploscompbiol/article?id=10.1371/journal.pcbi.1005633)。本轮是数据和指标准备，不是两篇论文的临床结果复现。', '']
    (output / '结论.md').write_text('\n'.join(lines), encoding='utf-8')
