"""One provisional observed-timing label and one held-out first-record score per person."""

from collections import Counter
from datetime import datetime, timezone
from itertools import combinations
import json
from pathlib import Path

import numpy as np

from ..provenance import file_hash
from .hospital_indicators import _json
from .hospital_proxy import (METHODS, TARGET_FIELDS, check_plan, circular_difference,
                             group_rows, held_out_scores, metrics, period_offset)


def validate_plan(plan, scoring):
    # PSEUDOCODE: reject any undeclared change to the prediction cutoff, endpoint or comparison thresholds.
    check_plan(scoring)
    expected = {
        'purpose': 'exploratory_person_level_observed_timing_variability_labels',
        'prediction_occasions': [1], 'calibration_occasions': [1], 'outcome_occasions': [2, 3, 4],
        'minimum_complete_outcome_occasions': 2,
        'label_aggregation': 'any_same_pair_sleep_midpoint_and_mean_meal_shift',
    }
    if any(plan.get(k) != v for k, v in expected.items()):
        raise ValueError('Unsupported person-level protocol; do not change the endpoint after seeing results.')


def clock_evidence(row):
    # PSEUDOCODE: accept measured clocks only; preserve missing/invalid evidence instead of imputing normal times.
    clocks = [row['values'].get(f) for f in TARGET_FIELDS]
    missing = [f for f, v in zip(TARGET_FIELDS, clocks)
               if v is None or not np.isfinite(v) or not 0 <= v < 1440]
    if missing:
        return None, missing
    start, end, *meals = clocks
    duration = (end - start) % 1440
    if duration == 0:
        return None, ['ambiguous_sleep_interval']
    return {'midpoint': (start + duration / 2) % 1440, 'meals': meals}, []


def annotate_person(records, scoring):
    """A zero means no qualifying change in the available observations, never confirmed clinical normality."""
    # PSEUDOCODE: use only later measurements -> compare complete pairs -> label observed variability with evidence.
    if [r['occasion'] for r in records] != [1, 2, 3, 4]:
        raise ValueError('Expect four ordered source occasions per person.')
    complete, missing, pairs = [], [], []
    for row in records[1:]:
        clocks, reasons = clock_evidence(row)
        if clocks is None:
            missing.append({'record_id': row['record_id'], 'reasons': reasons})
        else:
            complete.append((row, clocks))
    for (first, a), (second, b) in combinations(complete, 2):
        sleep = float(circular_difference(a['midpoint'], b['midpoint']))
        meal = float(np.mean([circular_difference(x, y) for x, y in zip(a['meals'], b['meals'])]))
        pairs.append({'records': [first['record_id'], second['record_id']],
                      'sleep_shift_min': sleep, 'meal_shift_min': meal, 'joint_shift_min': min(sleep, meal)})
    strongest = max(pairs, key=lambda p: p['joint_shift_min']) if pairs else None
    offsets = [period_offset(r) for r in records]
    dates_verified = all(d is not None for d in offsets) and all(a < b for a, b in zip(offsets, offsets[1:]))
    change = strongest['joint_shift_min'] if strongest else None
    return {'participant_id': records[0]['participant_id'],
        'annotation_kind': 'rule_inferred_observed_variability', 'clinical_label': None, 'clinical_onset': None,
        'joint_shift_min': change,
        'labels': {str(cut): int(change >= cut) if change is not None else None for cut in scoring['shift_minutes']},
        'complete_outcome_occasions': len(complete), 'missing_clock_evidence': missing,
        'dates_verified': dates_verified,
        'date_assumption': 'source_order_after_first_record' if not dates_verified else 'verified_month_day_order_year_unknown',
        'outcome_records': [r['record_id'] for r in records[1:]], 'pairs': pairs, 'strongest_pair': strongest}


def summarize_people(rows, scoring):
    # PSEUDOCODE: count each person once, report both classes and keep partial-evidence/date subsets explicit.
    ids = [r['participant_id'] for r in rows]
    if len(ids) != len(set(ids)):
        raise ValueError('Person-level evaluation must not count repeated records as extra people.')
    cutoff, q = scoring['primary_shift_minutes'], scoring['primary_alert_quantile']
    primary = {m: metrics(rows, m, cutoff, q) for m in METHODS}
    subsets = {'complete_three_outcome_records': [r for r in rows if r['complete_outcome_occasions'] == 3],
               'unambiguous_date_order': [r for r in rows if r['dates_verified']]}
    sensitivity = [{'shift_minutes': cut, 'alert_quantile': quantile, 'method': m,
                    **metrics(rows, m, cut, quantile)}
                   for cut in scoring['shift_minutes'] for quantile in scoring['alert_quantiles'] for m in METHODS]
    # Paired person resampling of fixed held-out predictions; it does not include model-refitting uncertainty.
    valid = [r for r in rows if r['joint_shift_min'] is not None and all(r['scores'][m] is not None for m in METHODS)]
    y = np.asarray([r['labels'][str(cutoff)] for r in valid])
    correct = {m: np.asarray([r['alerts'][m][str(q)] for r in valid]) == y for m in METHODS}
    rng = np.random.default_rng(scoring['seed'] + 2)
    samples = []
    for _ in range(scoring['bootstrap_replicates']):
        draw = rng.integers(len(valid), size=len(valid)) if valid else []
        if valid:
            samples.append(float(correct[METHODS[0]][draw].mean() - correct[METHODS[1]][draw].mean()))
    return {'people': len(rows), 'label_counts': dict(Counter(str(r['labels'][str(cutoff)]) for r in rows)),
        'complete_outcome_record_counts': dict(Counter(r['complete_outcome_occasions'] for r in rows)),
        'unambiguous_date_order_people': len(subsets['unambiguous_date_order']),
        'primary': primary, 'sensitivity': sensitivity,
        'subsets': {name: {'people': len(group), 'metrics': {m: metrics(group, m, cutoff, q) for m in METHODS}}
                    for name, group in subsets.items()},
        'paired_accuracy_difference': {
            'estimate': primary[METHODS[0]]['accuracy'] - primary[METHODS[1]]['accuracy'] if valid else None,
            'conditional_person_bootstrap_95_interval': np.quantile(samples, [.025, .975]).tolist() if samples else None,
            'replicates': len(samples), 'includes_refitting_uncertainty': False},
        'clinical_warning_accuracy': None, 'clinical_lead_time': None}


def write_workbook(rows, source_rows, scoring, path):
    # PSEUDOCODE: present one row per person with provisional label, held-out prediction, correctness and evidence.
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    grouped = group_rows(source_rows); cutoff = str(scoring['primary_shift_minutes']); q = str(scoring['primary_alert_quantile'])
    book = Workbook(); sheet = book.active; sheet.title = '60人暂定标签与预测'
    headers = ['人员编号', '暂定标签（1紊乱/0未见达线变化）', '标签性质', 'DNB判断（1提醒/0不提醒）', 'DNB是否判对',
               'DNB分数', '简单对照判断', '简单对照是否判对', '参考记录对的睡眠中点差（分钟）',
               '同一对记录的三餐平均变化（分钟）', '标注依据记录', '后续完整记录数', '日期先后能否核实',
               '缺失时刻说明', '原始工作表', '原始日期（记录顺序）']
    sheet.append(headers)
    for row in rows:
        label = row['labels'][cutoff]; pair = row['strongest_pair']
        alerts = [row['alerts'][m][q] for m in METHODS]
        correct = ['待核实' if label is None or a is None else '对' if label == a else '错' for a in alerts]
        original = grouped[row['participant_id']]
        sheet.append([row['participant_id'], label, '现有记录规则推定；非医院诊断', alerts[0], correct[0],
            row['scores'][METHODS[0]], alerts[1], correct[1], pair['sleep_shift_min'] if pair else None,
            pair['meal_shift_min'] if pair else None, ' / '.join(pair['records']) if pair else '',
            row['complete_outcome_occasions'], '可核实月日顺序' if row['dates_verified'] else '暂按原表次序；日期待核实',
            '; '.join(f'{m["record_id"]}: {", ".join(m["reasons"])}' for m in row['missing_clock_evidence']),
            original[0].get('sheet', ''), ' → '.join(str(r.get('raw', {}).get('date', {}).get('value', '')) for r in original)])
    sheet.freeze_panes = 'C2'; sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(color='FFFFFF', bold=True); cell.fill = PatternFill('solid', fgColor='234E70')
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    sheet.row_dimensions[1].height = 45
    for i in range(1, len(headers) + 1):
        sheet.column_dimensions[get_column_letter(i)].width = 22 if i not in (3, 11, 14, 16) else 40
    book.save(path); book.close()
    checked = load_workbook(path, read_only=True, data_only=True)
    values = list(checked.active.iter_rows(min_row=2, values_only=True)); checked.close()
    if len(values) != len(rows) or [(r[0], r[1], r[3]) for r in values] != [
        (r['participant_id'], r['labels'][cutoff], r['alerts'][METHODS[0]][q]) for r in rows]:
        raise ValueError('Workbook labels or predictions differ from the recorded experiment.')


def write_report(result, output):
    # PSEUDOCODE: state the person-level answer first and explain the provisional label in ordinary language.
    counts = result['label_counts']
    lines = ['# 按人标注与比对', '',
        f'共 {result["people"]} 人，每人一行。暂定紊乱（1）{counts.get("1", 0)} 人；未见达线变化（0）{counts.get("0", 0)} 人；证据不足 {counts.get("None", 0)} 人。', '',
        '标签规则：只看原表第 2—4 条记录；只要同一对记录中，睡眠中点相差至少 60 分钟、三餐时刻的平均差也至少 60 分钟，就记为 1；否则记为 0。至少需要两条完整记录。0 只表示已有记录未达到这条规则，不表示未记录的日子一定正常。', '',
        '模型只看第 1 条记录的 12 项数值指标，不看上述五个时刻字段；五折按人隔离参考、校准和测试。提醒门槛取独立校准人员首次分数的 90% 分位数；没有用答案调门槛。参考人员尚未确认稳定，因此这里是 sDNB 公式的探索评分。', '',
        '| 方法 | 判对/可比人数 | 正确率（暂定标签） | 命中紊乱 | 漏报 | 误报 | 平衡正确率 |',
        '|---|---:|---:|---:|---:|---:|---:|']
    for name, label in zip(METHODS, ('sDNB 探索评分', '同指标偏离对照')):
        m = result['primary'][name]
        pct = lambda v: f'{v:.2%}' if v is not None else '不可计算'
        lines.append(f'| {label} | {m.get("tp", 0)+m.get("tn", 0)}/{m["n"]} | {pct(m["accuracy"])} | '
                     f'{m.get("tp", 0)}/{m.get("positives", 0)} | {m.get("fn", 0)} | {m.get("fp", 0)} | {pct(m["balanced_accuracy"])} |')
    baseline = result['primary'][METHODS[0]]['always_no_change_accuracy']
    if baseline is not None:
        lines.extend(['', f'所有人都判 0 的正确率为 {baseline:.2%}，用于检查高正确率是否只是因为标签多数为 0。'])
    complete = result['complete_outcome_record_counts'].get(3, 0)
    lines += ['', f'{complete} 人有 3 条完整后续时刻记录，其余人的覆盖情况在 Excel 中逐人注明。'
        f'只有 {result["unambiguous_date_order_people"]} 人的全部月日先后可核实；其余暂按原表顺序。日期未补写，不能报告真实提前天数。', '',
        '**这些是现有记录的规则推定标签，不是医院确诊。** 一小时是沿用的探索尺度，变化也可能来自改善、周末或正常安排；这轮比对的是“后续记录中观察到的作息变动”，不能代替持续节律紊乱的真实预警验证。', '',
        '原始表、以前的结果均保留。participants.json 保存每对记录的依据；result.json 同时报告全部 30/60/90 分钟规则、三个提醒门槛、完整记录及日期可核实子集，没有选成绩最好的一格。', '',
        '陈洛南团队的 17 人流感实验使用实际症状/临床信息作为参照；我们当前的规则标签与其证据层级不同。[sDNB 原文](https://journals.plos.org/ploscompbiol/article?id=10.1371/journal.pcbi.1005633)', '']
    (output / '结论.md').write_text('\n'.join(lines), encoding='utf-8')


def run_participants(source, plan_path, output):
    # PSEUDOCODE: freeze protocol -> score first records -> independently annotate later evidence -> verify private artifacts.
    from . import hospital_proxy, reference_stability
    source, plan_path, output = Path(source), Path(plan_path), Path(output)
    plan = json.loads(plan_path.read_text(encoding='utf-8'))
    scoring_path = plan_path.parent / plan['scoring_plan']
    scoring = json.loads(scoring_path.read_text(encoding='utf-8')); validate_plan(plan, scoring)
    data = json.loads(source.read_text(encoding='utf-8')); source_digest = file_hash(source)
    grouped = group_rows(data['rows'])
    output.mkdir(parents=True, exist_ok=False)
    _json(output / 'protocol.json', {'created_at': datetime.now(timezone.utc).isoformat(),
        'plan': plan, 'scoring_plan': scoring, 'source': str(source.resolve()), 'source_sha256': source_digest,
        'original_workbook_sha256': data['source_sha256'],
        'plan_hashes': {str(p.resolve()): file_hash(p) for p in (plan_path, scoring_path)},
        'code_hashes': {str(p): file_hash(p) for p in (Path(__file__).resolve(), Path(hospital_proxy.__file__).resolve(),
                                                     Path(reference_stability.__file__).resolve())}})
    scores, lineage = held_out_scores(data['rows'], scoring, scored_occasions=(1,), calibration_occasions=(1,))
    _json(output / 'scores.json', {'predictions': scores, 'folds': lineage})
    rows = [{**score, **annotate_person(grouped[score['participant_id']], scoring)} for score in scores]
    result = summarize_people(rows, scoring)
    _json(output / 'participants.json', rows); _json(output / 'result.json', result)
    write_workbook(rows, data['rows'], scoring, output / '60人标签与预测.xlsx')
    write_report(result, output)
    if file_hash(source) != source_digest:
        raise ValueError('Source data changed during the experiment.')
    manifest = {'files': {p.name: file_hash(p) for p in sorted(output.iterdir()) if p.is_file()}}
    _json(output / 'manifest.json', manifest)
    if not all(file_hash(output / f) == h for f, h in manifest['files'].items()):
        raise ValueError('Saved artifacts failed hash verification.')
    return {'output': str(output.resolve()), 'people': result['people'], 'label_counts': result['label_counts'],
            'primary': result['primary']}
