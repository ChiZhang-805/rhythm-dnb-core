"""Audit the labeled workbook and publish a traceable experimental database."""

from collections import Counter, defaultdict
from datetime import datetime, timedelta
from itertools import zip_longest
import json
import math
from pathlib import Path
import shutil
import sqlite3

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..io.sources.hospital_workbook import BOUNDS
from ..provenance import file_hash, fingerprint
from ..text.schema import METRICS, FORMAL_CATEGORIES
from .hospital_participants import annotate_person


EXPERIMENT = '已标注实验数据'
HOSPITAL = '医院暂定标签数据'
OBJECTIVE = {
    'sleep_midpoint_h': '睡眠中点（小时）', 'sleep_duration_h': '睡眠时长（小时）',
    'meal_interval_cv': '餐间隔不规则程度', 'exercise_hour': '运动时刻（小时）',
    'exercise_minutes': '运动时长（分钟）', 'resting_hr_bpm': '静息心率',
    'screen_time_min': '屏幕使用（分钟）',
}
TEXT = {k: label for k, category, label, _ in METRICS if category in FORMAL_CATEGORIES}
HOSPITAL_FIELDS = {
    'age': '年龄', 'gender': '性别', 'sleep_duration_h': '睡眠时长（小时）',
    'time_in_bed': '卧床时长（小时）', 'nap_county': '小睡次数', 'nap_duration_min': '小睡时长（分钟）',
    'late_night_eating_flag': '是否夜宵', 'breakfast_kcal': '早餐热量（千卡）', 'lunch_kcal': '午餐热量（千卡）',
    'dinner_kcal': '晚餐热量（千卡）', 'water_ml': '饮水量（毫升）', 'stool_count': '排便次数',
    'smoking_state': '吸烟档位（0至5）', 'caffeine_mg': '咖啡因（毫克）', 'exercise_minutes': '运动时长（分钟）',
    'exercise_type': '运动类型', 'heart_rate_bpm': '心率（次/分）', 'resting_hr_bpm': '静息心率（次/分）',
    '_max_bpm': '最高心率（次/分）', 'skin_temp_c': '皮肤温度（℃）', 'spO2_pct': '血氧（%）',
    'spO2_min_pct': '最低血氧（%）', 'press_score': '压力（0至10）', 'general_mood': '心情（0至10）',
    'energy_score': '精力（0至10）', 'self_rate_state': '状态自评（1至5）', 'fatigue_score': '疲劳（0至10）',
    'social_interaction_min': '社交时长（分钟）', 'social_contact_count': '社交人数', 'ambient_temp_c': '环境温度（℃）',
    'noise_db': '噪声（分贝）', 'screen_time_min': '屏幕时长（分钟）', 'sleep_start_time': '入睡时刻',
    'sleep_end_time': '起床时刻', 'breakfast_time': '早餐时刻', 'lunch_time': '午餐时刻',
    'dinner_time': '晚餐时刻', 'caffeine_last_time': '末次咖啡因时刻', 'exercise_time_of_day': '运动时刻',
}
HOSPITAL_PANEL = ('sleep_duration_h', 'nap_duration_min', 'exercise_minutes', 'resting_hr_bpm',
                  'skin_temp_c', 'spO2_pct', 'press_score', 'general_mood', 'energy_score',
                  'fatigue_score', 'social_interaction_min', 'screen_time_min')
METHODS = {'personal_dnb': 'DNB', 'mean_deviation': '简单对照', 'history_dnb': '联合DNB',
           'history_control': '联合对照', 'rolling_dnb': '滚动DNB'}
ORIGINAL_METHODS = {'reference_robust_dnb': 'DNB', 'reference_robust_deviation': '简单对照',
                    'history_robust_dnb': '联合DNB', 'history_robust_control': '联合对照'}
FUTURE_REASONS = {
    'baseline_only': '稳定基线期，不参与未来事件判分',
    'first_event_already_started': '事件已经开始，不再预测首次发生',
    'insufficient_future_confirmation_coverage': '末尾随访不足以判定未来7天',
    'complete_simulated_followup': '已有封存参考答案',
}


def require(value, message):
    if not value:
        raise ValueError(message)


def equal(a, b):
    if a is None or b is None:
        return a is b
    if type(a) in (int, float) and type(b) in (int, float):
        return math.isfinite(a) and math.isfinite(b) and math.isclose(a, b, rel_tol=1e-13, abs_tol=1e-11)
    return a == b


def instant(value):
    stamp = datetime.fromisoformat(value)
    require(stamp.tzinfo is not None, '时间戳必须带时区')
    return stamp


def simulated_clocks(midpoint_h, duration_h):
    """Reuse the saved simulator's clock formula; never infer measured hospital clocks."""
    require(type(midpoint_h) in (int, float) and math.isfinite(midpoint_h) and 0 <= midpoint_h < 24, '无效睡眠中点')
    require(type(duration_h) in (int, float) and math.isfinite(duration_h) and 0 < duration_h < 24, '无效睡眠时长')
    result = []
    for sign in (-1, 1):
        microseconds = round(((midpoint_h + sign * duration_h / 2) % 24) * 3_600_000_000) % 86_400_000_000
        seconds, fraction = divmod(microseconds, 1_000_000)
        hour, rest = divmod(seconds, 3600)
        minute, second = divmod(rest, 60)
        result.append(f'{hour:02d}:{minute:02d}:{second:02d}.{fraction:06d}')
    return tuple(result)


def hospital_cross_checks(values):
    # 仅标出逻辑矛盾；不能靠猜测交换数值或删去小数点。
    checks = [('sleep_duration_h', 'time_in_bed', '睡眠时长超过卧床时长'),
              ('spO2_min_pct', 'spO2_pct', '最低血氧高于平均血氧'),
              ('resting_hr_bpm', '_max_bpm', '静息心率高于最高心率'),
              ('heart_rate_bpm', '_max_bpm', '平均心率高于最高心率')]
    return [(a, b, message) for a, b, message in checks
            if values.get(a) is not None and values.get(b) is not None and values[a] > values[b] + 1e-10]


def read_json(path):
    return json.loads(Path(path).read_text(encoding='utf-8'))


def save_json(path, obj):
    with Path(path).open('x', encoding='utf-8') as f:
        json.dump(obj, f, ensure_ascii=False, indent=2, allow_nan=False)
        f.write('\n')


def read_book(path):
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        require(book.sheetnames == [EXPERIMENT, HOSPITAL], 'Unexpected workbook sheets')
        output = {}
        for sheet in book:
            it = sheet.iter_rows(values_only=True)
            header = tuple(next(it))
            require(len(header) == len(set(header)), 'Duplicate column names')
            rows = []
            for values in it:
                require(len(values) <= len(header), 'Unexpected trailing cells')
                values = tuple(values) + (None,) * (len(header) - len(values))
                for value in values:
                    require(not isinstance(value, float) or math.isfinite(value), 'Nonfinite source value')
                rows.append(dict(zip(header, values)))
            require(len({r['记录编号'] for r in rows}) == len(rows), 'Duplicate record IDs')
            output[sheet.title] = {'headers': list(header), 'rows': rows}
        return output
    finally:
        book.close()


def restore(row, column, expected, changes, source):
    # 只补空格；已有值不相符时停止，不能静默覆盖历史实验。
    value = row[column]
    if value is None and expected is not None:
        row[column] = expected
        changes.append({'record_id': row['记录编号'], 'column': column, 'old': None, 'new': expected, 'source': source})
    else:
        require(equal(value, expected), f'Source disagreement: {row["记录编号"]} / {column}')


def load_source(database):
    with sqlite3.connect(Path(database).resolve().as_uri() + '?mode=ro', uri=True) as db:
        db.row_factory = sqlite3.Row
        db.execute('BEGIN')
        require(db.execute('PRAGMA quick_check').fetchone()[0] == 'ok', 'Source database integrity failed')
        originals = {r['record_id']: dict(r) for r in db.execute('''
            SELECT o.*, r.payload AS raw, a.reference_event_in_followup
            FROM observations o JOIN raw_inputs r USING(record_id)
            JOIN research_outcome_annotations a USING(record_id)
            WHERE a.reference_event_in_followup IN (0,1)''')}
        followup = {r['record_id']: dict(r) for r in db.execute('''
            SELECT d.*, s.parent_participant_id, s.split, s.event_in_followup,
                   s.first_onset_date,s.first_confirmed_date,s.followup_end_date,u.created_at AS annotation_created_at
            FROM research_followup_days d JOIN research_followup_sequences s USING(sequence_id)
            JOIN research_followup_runs u USING(run_id)''')}
        forecasts = defaultdict(dict)
        for table, mapping in [('research_warning_annotations', ORIGINAL_METHODS), ('research_expanded_forecasts', METHODS)]:
            for r in db.execute('SELECT * FROM ' + table):
                if r['method'] in mapping:
                    forecasts[r['record_id']][mapping[r['method']]] = dict(r)
        text = defaultdict(dict)
        for r in db.execute('''SELECT b.record_id,s.estimates FROM research_text_bindings b
                           JOIN research_text_scores s USING(run_id,task_id)'''):
            text[r['record_id']].update(json.loads(r['estimates']))
    return originals, followup, forecasts, text


def audit_experiment(rows, source, changes):
    originals, followups, forecasts, text = source
    require({r['记录编号'] for r in rows} == set(originals) | set(followups), 'Labeled record coverage changed')
    quality, people, timelines = {}, {}, defaultdict(list)
    for row in rows:
        rid, person = row['记录编号'], row['人物编号']
        followup = rid in followups
        original = followups[rid] if followup else originals[rid]
        group = 'SIMULATED-FOLLOWUP' if followup else 'SIMULATED-RHYTHM'
        require(row['数据分组'] == group, 'Cohort mismatch')
        sequence = original['sequence_id'] if followup else original['participant_id']
        require(person == (original['parent_participant_id'] if followup else sequence) and row['序列编号'] == sequence, 'Identity mismatch')
        label = original['event_in_followup'] if followup else original['reference_event_in_followup']
        require(row['随访期间是否紊乱'] == label and label in (0, 1), 'Outcome mismatch')
        require(row['标签文字'] == ('紊乱' if label else '未紊乱'), 'Label text mismatch')
        role = original['split']
        require(row['训练或测试分组'] == role, 'Frozen split changed')
        identity = (label, role, sequence)
        require(person not in people or people[person] == identity, 'Person labels or splits conflict')
        people[person] = identity
        measured = original['simulated_at'] if followup else original['observed_at']
        require(instant(row['记录时间']) == instant(measured), 'Measurement time mismatch')
        require(instant(row['预测时点']) >= instant(measured), 'Prediction precedes available inputs')
        timelines[sequence].append(instant(measured))
        values = json.loads(original['features_json']) if followup else dict(original)
        if not followup:
            raw = json.loads(original['raw'])
            a, b = instant(raw['sleep_start_time']), instant(raw['sleep_end_time'])
            require(a < b <= instant(measured) and b-a <= timedelta(days=1), 'Invalid sleep interval')
            middle = a + (b-a)/2
            values['sleep_midpoint_h'] = middle.hour + middle.minute/60 + middle.second/3600 + middle.microsecond/3.6e9
            for field, column in [('sleep_start_time', '入睡时间'), ('sleep_end_time', '起床时间')]:
                restore(row, column, raw[field], changes, 'raw_inputs.' + field)
            for category, column in [('emotion', '情绪文本'), ('sleep', '睡眠文本'), ('diet', '饮食文本'), ('social', '社交文本')]:
                restore(row, column, raw[category + '_description_text'], changes, 'raw_inputs.' + category)
            for key, name in TEXT.items():
                restore(row, '参考·' + name, original['text_' + key], changes, 'observations.text_' + key)
                restore(row, '模型·' + name, text[rid][key], changes, 'research_text_scores.' + key)
        for key, name in OBJECTIVE.items():
            restore(row, name, values[key], changes, 'research_followup_days.features_json' if followup else 'observations/aware_sleep_midpoint')
            v = row[name]
            require(type(v) in (int, float) and math.isfinite(v) and v >= 0, 'Missing/invalid objective feature')
            upper = 24 if key in ('sleep_midpoint_h', 'exercise_hour', 'sleep_duration_h') else 1440 if key.endswith('minutes') or key == 'screen_time_min' else math.inf
            require(v <= upper and (key not in ('sleep_midpoint_h', 'exercise_hour') or v < upper), 'Invalid objective range')
        state = original['regime_state'] if followup else original['rhythm_state_gt']
        restore(row, '当日是否紊乱', state, changes, 'existing_day_state')
        onset = original['first_onset_date'] if followup else original['event_onset_at']
        end = original['followup_end_date'] if followup else original['followup_end_at']
        restore(row, '紊乱开始时间', onset, changes, 'existing_event_onset')
        restore(row, '随访结束时间', end, changes, 'existing_followup_end')
        require((onset is not None) == bool(label), 'Event/onset inconsistency')
        require(measured[:10] <= end[:10] and (onset is None or onset[:10] <= end[:10]), 'Event outside follow-up')
        expected_state = int(onset is not None and measured[:10] >= onset[:10])
        require(state == expected_state, 'Daily state disagrees with authored first event')
        notes, exceptions = [], {}
        if followup:
            clock_a, clock_b = simulated_clocks(values['sleep_midpoint_h'], values['sleep_duration_h'])
            restore(row, '入睡时间', clock_a, changes, 'saved_simulator_midpoint_minus_half_duration; clock_only')
            restore(row, '起床时间', clock_b, changes, 'saved_simulator_midpoint_plus_half_duration; clock_only')
            restore(row, '参考答案可获得时间', original['annotation_created_at'], changes, 'research_followup_runs.created_at; retrospective_annotation_time')
            restore(row, '紊乱确认时间', original['first_confirmed_date'], changes, 'research_followup_sequences.first_confirmed_date')
            restore(row, '未来7天是否开始紊乱', original['future_event_7d'], changes, 'research_followup_days.future_event_7d')
            future_status = FUTURE_REASONS[original['future_label_status']]
            notes.append('无文本采集；睡眠时刻按模拟规则反算（不含日期）')
            exceptions.update({'入睡时间': 'derived_simulated_clock', '起床时间': 'derived_simulated_clock', '参考答案可获得时间': 'restored_annotation_creation_time'})
            for col in row:
                if col.endswith('文本') or col.startswith(('参考·', '模型·')):
                    require(row[col] is None, 'Follow-up text would be newly invented')
                    exceptions[col] = 'not_collected_in_numeric_followup'
        else:
            restore(row, '参考答案可获得时间', original['label_observed_at'], changes, 'observations.label_observed_at')
            future_status = '已有封存参考答案' if row['未来7天是否开始紊乱'] is not None else '此行未纳入原封存7天预测评价'
            if label:
                exceptions['紊乱确认时间'] = 'not_defined_in_original_scenario'
                notes.append('原情景未单列确认时刻')
        for method in METHODS.values():
            saved = forecasts.get(rid, {}).get(method)
            for suffix, field in [('分数', 'score'), ('预警', 'warning')]:
                col = method + suffix
                restore(row, col, saved[field] if saved else None, changes, 'saved_forecast.' + method)
                if saved is None:
                    exceptions[col] = 'not_run_in_frozen_experiment'
        method_basis = '个人稳定日参考' if followup else '稳定参考人群'
        restore(row, 'DNB参考方式', method_basis if row['DNB分数'] is not None else None, changes, 'saved_forecast_reference_basis')
        if not label:
            exceptions.update({'紊乱开始时间': 'not_applicable_no_event', '紊乱确认时间': 'not_applicable_no_event'})
        if row['未来7天是否开始紊乱'] is None:
            exceptions['未来7天是否开始紊乱'] = future_status
        if row['DNB参考方式'] is None:
            exceptions['DNB参考方式'] = 'not_run_in_frozen_experiment'
        for key, value in row.items():
            if value is None:
                require(key in exceptions, f'Unclassified blank: {rid}/{key}')
            elif key.startswith(('参考·', '模型·')):
                require(type(value) in (int, float) and 0 <= value <= 100, 'Text score out of range')
        row.update({'输入完整性': '7项量化指标完整', '7天标签状态': future_status,
                    '结果状态': '已有封存预测' if row['DNB分数'] is not None else '原实验未在此行输出预警',
                    '数据说明': '；'.join(notes) or '原始模拟文本与分数保留'})
        quality[rid] = {'exceptions': exceptions, 'objective_ready': True, 'text_ready': not followup,
                        'notes': notes, 'label_kind': 'simulation'}
    for seq, stamps in timelines.items():
        ordered = sorted(stamps)
        require(len(set(ordered)) == len(ordered), 'Duplicate sequence/day')
        require(all(b-a == timedelta(days=1) for a, b in zip(ordered, ordered[1:])), 'Daily sequence has gaps')
        require(len(ordered) == (84 if seq.startswith('FOLLOWUP-') else 14), 'Sequence length changed')
    return quality, people


def audit_hospital(rows, normalized, participants, scoring, changes):
    sources = {r['record_id']: r for r in normalized['rows']}
    people = {p['participant_id']: p for p in participants}
    require({r['记录编号'] for r in rows} == set(sources), 'Hospital record coverage mismatch')
    grouped = defaultdict(list)
    for r in sources.values():
        grouped[r['participant_id']].append(r)
    # 重新用同一已冻结规则核验标签，而不根据DNB输出改标签。
    for person, group in grouped.items():
        actual = annotate_person(sorted(group, key=lambda r: r['occasion']), scoring)
        require(actual['labels'] == people[person]['labels'], 'Existing hospital labels no longer reproducible')
    quality, issues = {}, []
    for row in rows:
        rid = row['记录编号']; source = sources[rid]; p = people[source['participant_id']]
        require(row['人物编号'] == source['participant_id'] and row['记录次序'] == source['occasion'], 'Hospital identity mismatch')
        require(row['人员暂定标签（1紊乱/0未紊乱）'] == p['labels']['60'], 'Hospital label changed')
        require(row['标签文字'] == ('紊乱' if p['labels']['60'] else '未紊乱'), 'Hospital text label mismatch')
        require(str(source['raw']['date']['value']) == row['原始日期'], 'Raw hospital date changed')
        restore(row, '可识别月日', source['date']['month_day'], changes, 'verified_hospital_partial_date')
        restore(row, '完整日期', source['date']['calendar_date'], changes, 'verified_hospital_date')
        exceptions, notes = {}, []
        if row['完整日期'] is None:
            exceptions['完整日期'] = 'year_not_supplied'
        if row['可识别月日'] is None:
            exceptions['可识别月日'] = source['date']['status']
        for key, col in HOSPITAL_FIELDS.items():
            value = source['values'][key]
            if key.endswith('_time') or key == 'exercise_time_of_day':
                if value is not None:
                    seconds = round(value * 60); hour, rest = divmod(seconds, 3600); minute, sec = divmod(rest, 60)
                    value = f'{hour:02d}:{minute:02d}:{sec:02d}'
            restore(row, col, value, changes, 'normalized_hospital_source.' + key)
            state = source['states'][key]
            if state != 'observed':
                exceptions[col] = state
                if state != 'not_applicable':
                    notes.append(col + '待核实')
                    issues.append({'record_id': rid, 'column': col, 'reason': state,
                                   'source_sheet': source['sheet'], 'source_cell': source['raw'][key]['cell'],
                                   'raw_value': source['raw'][key]['value']})
            if type(value) in (int, float) and key in BOUNDS:
                lo, hi = BOUNDS[key]; require(lo <= value <= hi, 'Invalid accepted hospital value')
        for a, b, reason in hospital_cross_checks(source['values']):
            notes.append(reason + '，保留原值待核实')
            for key in (a, b):
                exceptions[HOSPITAL_FIELDS[key]] = 'cross_field_conflict_pending_confirmation'
            issues.append({'record_id': rid, 'column': HOSPITAL_FIELDS[a] + '/' + HOSPITAL_FIELDS[b],
                           'reason': reason, 'source_sheet': source['sheet'],
                           'source_cell': source['raw'][a]['cell'] + '/' + source['raw'][b]['cell'],
                           'raw_value': [source['values'][a], source['values'][b]]})
        for prefix, method in [('sDNB', 'sdnb_exploratory'), ('对照', 'mean_deviation')]:
            for suffix, expected in [('分数', p['scores'][method]), ('预警', p['alerts'][method]['0.9'])]:
                col = '首条记录' + prefix + suffix
                restore(row, col, expected if source['occasion'] == 1 else None, changes, 'frozen_hospital_person_prediction')
                if source['occasion'] != 1:
                    exceptions[col] = 'outcome_record_not_prediction_record'
        for key, value in row.items():
            if value is None:
                require(key in exceptions, f'Unclassified hospital blank: {rid}/{key}')
        ready = all(source['values'][k] is not None and HOSPITAL_FIELDS[k] not in exceptions for k in HOSPITAL_PANEL)
        row.update({'输入完整性': '12项数值完整且无已知冲突' if ready else '部分指标待核实',
                    '数据说明': '；'.join(notes) or '无需补造数值；未发生行为的时刻不适用'})
        quality[rid] = {'exceptions': exceptions, 'objective_ready': ready, 'text_ready': False,
                        'notes': notes, 'label_kind': 'hospital_rule_provisional'}
    return quality, issues


def write_book(tables, quality, path):
    book = Workbook(write_only=True)
    for title, table in tables.items():
        sheet = book.create_sheet(title); sheet.freeze_panes = 'E2'
        sheet.sheet_view.showGridLines = False; sheet.row_dimensions[1].height = 44
        headers, rows = table['headers'], table['rows']
        sheet.auto_filter.ref = f'A1:{get_column_letter(len(headers))}{len(rows)+1}'
        cells = []
        for i, h in enumerate(headers, 1):
            sheet.column_dimensions[get_column_letter(i)].width = 38 if h.endswith('文本') or h == '数据说明' else 23
            cell = WriteOnlyCell(sheet, value=h)
            cell.font = Font(name='Microsoft YaHei', bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='22685C' if i in (3, 4) else '263B52')
            cell.alignment = Alignment(wrap_text=True, vertical='center')
            if i == 3:
                cell.comment = Comment('这是人员随访标签，不是每条记录当日的状态，也不是DNB输入。医院表为规则暂定标签。', '字段说明')
            if h == '参考答案可获得时间':
                cell.comment = Comment('扩展模拟记录填实际标注创建时间；与模拟事件日期不同，不可作为预测输入。', '字段说明')
            if h in ('入睡时间', '起床时间'):
                cell.comment = Comment('原模拟记录保留时间戳；扩展记录按原生成规则反算时刻，不补造真实日期。', '字段说明')
            cells.append(cell)
        sheet.append(cells)
        for row in rows:
            cells = []
            for header in headers:
                value = row[header]; cell = WriteOnlyCell(sheet, value=value)
                if isinstance(value, str):
                    require(len(value) <= 32767, 'Excel text truncation')
                    cell.data_type = 's'
                if isinstance(value, float):
                    cell.number_format = '0.00'
                state = quality[row['记录编号']]['exceptions'].get(header)
                if title == HOSPITAL and state:
                    cell.comment = Comment('空格或疑点原因：' + state + '。不自动填零、不按前后记录猜测。', '逐行核对')
                    if state != 'not_applicable' and state != 'outcome_record_not_prediction_record':
                        cell.fill = PatternFill('solid', fgColor='FFF1D6')
                cells.append(cell)
            sheet.append(cells)
    book.save(path)
    verified = read_book(path)
    for name, table in tables.items():
        require(verified[name]['headers'] == table['headers'], 'Export schema mismatch')
        for expected, actual in zip_longest(table['rows'], verified[name]['rows']):
            require(expected is not None and actual is not None, 'Export row count mismatch')
            require(all(equal(expected[k], actual[k]) for k in table['headers']), 'Export cell mismatch')


def quote(name):
    return '"' + name.replace('"', '""') + '"'


def validate_input_columns(view, names):
    # 使用精确白名单；sleep_onset_difficulty是输入症状，不是事件发生时间。
    identity = {'record_id', 'participant_id', 'sequence_id', 'dataset', 'split', 'observed_at', 'issued_at'}
    expected = {
        'dnb_objective_inputs': identity | set(OBJECTIVE),
        'dnb_text_inputs': identity | set(OBJECTIVE) | {'text_' + k for k in TEXT},
        'hospital_first_inputs': {'record_id', 'participant_id', 'values_verified'} | set(HOSPITAL_PANEL),
    }
    require(view in expected and len(names) == len(set(names)) and set(names) == expected[view],
            'Unexpected field or target leakage in input view')


def write_database(tables, quality, changes, sources, path):
    require(not path.exists(), 'Never overwrite an existing database')
    with sqlite3.connect(path) as db:
        db.execute('PRAGMA foreign_keys=ON')
        db.executescript('''
            CREATE TABLE metadata(key TEXT PRIMARY KEY, value_json TEXT NOT NULL);
            CREATE TABLE participant_outcomes(participant_id TEXT PRIMARY KEY, sequence_id TEXT NOT NULL,
                dataset TEXT NOT NULL, split TEXT, label INTEGER NOT NULL CHECK(label IN (0,1)),
                label_kind TEXT NOT NULL, onset TEXT, clinical_label INTEGER CHECK(clinical_label IN (0,1)));
            CREATE TABLE quality_profiles(profile_id TEXT PRIMARY KEY, quality_json TEXT NOT NULL);
            CREATE TABLE record_audit(record_id TEXT PRIMARY KEY, participant_id TEXT NOT NULL REFERENCES participant_outcomes,
                profile_id TEXT NOT NULL REFERENCES quality_profiles, objective_ready INTEGER NOT NULL, text_ready INTEGER NOT NULL);
            CREATE TABLE changes(record_id TEXT NOT NULL REFERENCES record_audit, field TEXT NOT NULL,
                old_value_json TEXT NOT NULL, new_value_json TEXT NOT NULL, evidence TEXT NOT NULL,
                PRIMARY KEY(record_id,field));
            CREATE TABLE columns(table_name TEXT NOT NULL, column_name TEXT NOT NULL, role TEXT NOT NULL,
                PRIMARY KEY(table_name,column_name));
        ''')
        people, all_ids = {}, set()
        for title, table in tables.items():
            name = 'experiment_records' if title == EXPERIMENT else 'hospital_records'
            headers, rows = table['headers'], table['rows']
            types = {}
            for h in headers:
                values = [r[h] for r in rows if r[h] is not None]
                types[h] = 'REAL' if values and all(type(x) in (int, float) for x in values) else 'TEXT'
            db.execute(f'CREATE TABLE {name} (' + ','.join(quote(h) + ' ' + types[h] + (' PRIMARY KEY' if h == '记录编号' else '') for h in headers) + ')')
            db.executemany(f'INSERT INTO {name} VALUES (' + ','.join('?' for _ in headers) + ')', [[r[h] for h in headers] for r in rows])
            input_names = set(OBJECTIVE.values()) if title == EXPERIMENT else set(HOSPITAL_FIELDS.values())
            for h in headers:
                role = 'input' if h in input_names or h.endswith('文本') or h.startswith('模型·') else 'reference_only' if h.startswith('参考·') else 'output_or_metadata'
                db.execute('INSERT INTO columns VALUES (?,?,?)', (name, h, role))
            for r in rows:
                rid, person = r['记录编号'], r['人物编号']; require(rid not in all_ids, 'Cross-table duplicate record')
                all_ids.add(rid)
                q = quality[rid]
                values = (person, r['序列编号'] if title == EXPERIMENT else person,
                          r['数据分组'] if title == EXPERIMENT else 'HOSPITAL',
                          r['训练或测试分组'] if title == EXPERIMENT else None,
                          r['随访期间是否紊乱'] if title == EXPERIMENT else r['人员暂定标签（1紊乱/0未紊乱）'],
                          q['label_kind'], r.get('紊乱开始时间'), None)
                require(person not in people or people[person] == values, 'Person outcome changed between records')
                if person not in people:
                    people[person] = values
                    db.execute('INSERT INTO participant_outcomes VALUES (?,?,?,?,?,?,?,?)', values)
                profile = fingerprint(q)
                db.execute('INSERT OR IGNORE INTO quality_profiles VALUES (?,?)', (profile, json.dumps(q, ensure_ascii=False, allow_nan=False)))
                db.execute('INSERT INTO record_audit VALUES (?,?,?,?,?)', (rid, person, profile, int(q['objective_ready']), int(q['text_ready'])))
        db.executemany('INSERT INTO changes VALUES (?,?,?,?,?)', [(c['record_id'], c['column'], json.dumps(c['old']), json.dumps(c['new'], ensure_ascii=False), c['source']) for c in changes])
        for key, value in sources.items():
            db.execute('INSERT INTO metadata VALUES (?,?)', (key, json.dumps(value, ensure_ascii=False)))
        identities = 'e."记录编号" AS record_id,e."人物编号" AS participant_id,e."序列编号" AS sequence_id,e."数据分组" AS dataset,e."训练或测试分组" AS split,e."记录时间" AS observed_at,e."预测时点" AS issued_at'
        columns = ','.join('e.' + quote(label) + ' AS ' + quote(key) for key, label in OBJECTIVE.items())
        db.execute(f'CREATE VIEW dnb_objective_inputs AS SELECT {identities},{columns} FROM experiment_records e')
        model_cols = ','.join('e.' + quote('模型·'+label) + ' AS ' + quote('text_'+key) for key, label in TEXT.items())
        db.execute(f'CREATE VIEW dnb_text_inputs AS SELECT {identities},{columns},{model_cols} FROM experiment_records e JOIN record_audit a ON e."记录编号"=a.record_id WHERE a.text_ready=1')
        columns = ','.join('h.' + quote(HOSPITAL_FIELDS[k]) + ' AS ' + quote(k) for k in HOSPITAL_PANEL)
        db.execute(f'''CREATE VIEW hospital_first_inputs AS SELECT h."记录编号" AS record_id,h."人物编号" AS participant_id,
                    {columns},a.objective_ready AS values_verified FROM hospital_records h JOIN record_audit a
                    ON h."记录编号"=a.record_id WHERE h."记录次序"=1''')
        db.commit()
        require(db.execute('PRAGMA integrity_check').fetchone()[0] == 'ok' and not db.execute('PRAGMA foreign_key_check').fetchall(), 'Database integrity failed')
        require(db.execute('SELECT count(*) FROM record_audit').fetchone()[0] == 23340, 'Database row coverage failed')
        require(db.execute('SELECT count(*) FROM participant_outcomes').fetchone()[0] == 485, 'Database people coverage failed')
        require(db.execute('SELECT count(*) FROM participant_outcomes WHERE clinical_label IS NOT NULL').fetchone()[0] == 0, 'Unverified clinical label inserted')
        for name, table in [('experiment_records', tables[EXPERIMENT]), ('hospital_records', tables[HOSPITAL])]:
            expected = {r['记录编号']: r for r in table['rows']}
            cursor = db.execute('SELECT * FROM ' + name); headers = [c[0] for c in cursor.description]
            for row in cursor:
                actual = dict(zip(headers, row)); require(all(equal(actual[k], expected[actual['记录编号']][k]) for k in headers), 'SQLite cell mismatch')
        for view in ('dnb_objective_inputs', 'dnb_text_inputs', 'hospital_first_inputs'):
            names = [r[1] for r in db.execute('PRAGMA table_info(' + view + ')')]
            validate_input_columns(view, names)


def prepare(workbook, database, hospital_normalized, hospital_people, hospital_protocol, output):
    workbook, database, output = Path(workbook).resolve(), Path(database).resolve(), Path(output).resolve()
    destination = workbook.with_suffix('.sqlite')
    require(not output.exists() and not destination.exists(), 'Output already exists; preserve it')
    paths = [workbook, database, Path(hospital_normalized).resolve(), Path(hospital_people).resolve(), Path(hospital_protocol).resolve()]
    hashes = {str(p): file_hash(p) for p in paths}
    output.mkdir(parents=True, exist_ok=False)
    backup = output / 'labeled_data.before.xlsx'
    shutil.copy2(workbook, backup)
    require(file_hash(backup) == hashes[str(workbook)], 'Backup verification failed')
    save_json(output / 'plan.json', {'source_hashes': hashes, 'goal': '逐行核验、可追溯补全、形成实验数据库',
        'allowed_fills': ['已有源数据的精确回填', '按已保存模拟规则反算睡眠时刻', '回填标注创建时间'],
        'never_fill': ['无原文的情绪等文本', '真实缺测数值', '不明年份或月日', '未知临床诊断', '未执行的预警结果'],
        'preserve': '原始库、标签、划分、指标值、已有预测均不修改；来源与临床标签分离。',
        'code_sha256': file_hash(__file__)})
    tables = read_book(workbook); source = load_source(database)
    require((len(tables[EXPERIMENT]['rows']), len(tables[HOSPITAL]['rows'])) == (23100, 240), 'Unexpected source size')
    original_snapshot = {r['记录编号']: dict(r) for t in tables.values() for r in t['rows']}
    changes = []
    q1, persons = audit_experiment(tables[EXPERIMENT]['rows'], source, changes)
    print('23,100 experimental records verified against the source database.', flush=True)
    q2, issues = audit_hospital(tables[HOSPITAL]['rows'], read_json(hospital_normalized), read_json(hospital_people), read_json(hospital_protocol)['scoring_plan'], changes)
    quality = q1 | q2
    for title, extra in [(EXPERIMENT, ['输入完整性','7天标签状态','结果状态','数据说明']), (HOSPITAL,['输入完整性','数据说明'])]:
        tables[title]['headers'].extend(extra)
    # 每个原有非空单元格都必须保持原值，包括疑点、标签、划分和旧预测。
    for table in tables.values():
        for row in table['rows']:
            for col, value in original_snapshot[row['记录编号']].items():
                if value is not None:
                    require(equal(row[col], value), 'An existing value was altered')
    stats = {}
    for title, table in tables.items():
        missing = Counter(); before_missing = Counter(); reasons = Counter()
        for row in table['rows']:
            for col, val in original_snapshot[row['记录编号']].items():
                before_missing[col] += int(val is None)
                missing[col] += int(row[col] is None)
            reasons.update(quality[row['记录编号']]['exceptions'].values())
        stats[title] = {'rows': len(table['rows']), 'missing_before': dict(before_missing), 'missing_after': dict(missing), 'cell_reason_counts': dict(reasons)}
    staged_book, staged_db = output / workbook.name, output / destination.name
    write_book(tables, quality, staged_book)
    print('Excel roundtrip verified; writing and checking the SQLite database.', flush=True)
    write_database(tables, quality, changes, {'source_hashes': hashes, 'input_workbook': str(workbook), 'scope': 'Simulation and provisional hospital-rule outcomes; no independently confirmed clinical endpoint'}, staged_db)
    require(all(file_hash(p) == hashes[str(p)] for p in paths), 'Source changed during processing')
    change_counts = Counter(c['column'] for c in changes)
    report = {'source_hashes': hashes, 'rows_checked': 23340, 'people_or_trajectories': 485,
              'existing_nonempty_cells_preserved': True, 'all_blank_cells_classified': True,
              'new_clinical_labels': 0, 'frozen_splits_preserved': True,
              'filled_cells': len(changes), 'filled_by_column': dict(change_counts), 'tables': stats,
              'hospital_source_issues': issues, 'hospital_calendar_year_unknown_records': 240,
              'hospital_ambiguous_month_day_records': sum(r['可识别月日'] is None for r in tables[HOSPITAL]['rows']),
              'hospital_verified_numeric_rows': sum(q['objective_ready'] for q in q2.values()),
              'objective_input_rows': len(q1), 'text_input_rows': sum(q['text_ready'] for q in q1.values()),
              'workbook_sha256': file_hash(staged_book), 'database_sha256': file_hash(staged_db),
              'backup': str(backup), 'database': str(destination), 'workbook': str(workbook)}
    save_json(output / 'verification.json', report)
    lines = ['# 数据检查结果', '',
             '已逐行核对 23,340 条记录、485 个实验人物／序列，保留全部原有数值、标签、划分和历史结果。', '',
             f'补入 {len(changes):,} 个空格：' + '；'.join(f'{k} {v:,} 格' for k,v in change_counts.items()) + '。', '',
             '23,100 条模拟记录的 7 项量化指标均完整；其中 2,520 条有文本及模型分数。扩展随访没有采集文本，不能补写文字后冒充独立输入。', '',
             f'医院数据保留 60 人、240 条记录；{report["hospital_verified_numeric_rows"]} 条的既定 12 项数值完整且没有本次发现的冲突。年份均待核实，66 条月日不清；另有少量缺测、越界和 1 条血氧矛盾，原值与原因均可查。', '',
             'Excel仍为两张数据表。SQLite的 dnb_objective_inputs、dnb_text_inputs、hospital_first_inputs 只提供相应实验输入；人员结局在 participant_outcomes。不得把后续记录当成人员预测时已知的信息。', '',
             '这是统一的实验数据源：425 个模拟人物／序列与60名医院暂定标签人员分别使用。模拟结局和规则推定不是独立临床结局；这次整理不改变这一点。', '']
    (output / '检查结果.md').write_text('\n'.join(lines), encoding='utf-8')
    # 先发布新库，再用同目录原子替换更新Excel。锁定时保留可核验的新文件，不破坏原件。
    with staged_db.open('rb') as src, destination.open('xb') as dst:
        shutil.copyfileobj(src, dst)
    require(file_hash(destination) == report['database_sha256'], 'Database publication failed')
    pending = workbook.with_name('.labeled_data.checked.xlsx')
    require(not pending.exists(), 'A previous staged replacement exists')
    with staged_book.open('rb') as src, pending.open('xb') as dst:
        shutil.copyfileobj(src, dst)
    require(file_hash(pending) == report['workbook_sha256'], 'Staged Excel publication failed')
    try:
        pending.replace(workbook)
    except PermissionError:
        save_json(output / 'installed.json', {'workbook_locked': True, 'checked_workbook': str(pending), 'database': str(destination)})
        return {**report, 'workbook_locked': True, 'checked_workbook': str(pending)}
    require(file_hash(workbook) == report['workbook_sha256'], 'Final Excel verification failed')
    save_json(output / 'installed.json', {'workbook': str(workbook), 'database': str(destination),
        'workbook_sha256': report['workbook_sha256'], 'database_sha256': report['database_sha256']})
    return report
