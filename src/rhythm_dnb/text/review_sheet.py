"""Editable blind-review workbooks with exact source checks and explicit human attestation."""

from copy import deepcopy
from pathlib import Path
from ..provenance import file_hash
from .schema import CATEGORIES, METRICS
from .labels import STATES, validate_labels
from .annotations import validate_scope_targets

HEADERS = ('盲编号', '类别', '原文', '指标编号', '指标', '标注者', '证据状态', '分数', '对应时段', '依据原文', '排除原文', '排除原因', '本人已独立复核')


def export_review_sheet(packet, path):
    # PSEUDOCODE: expose original text and empty answers in a spreadsheet while withholding old scores and other raters.
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.worksheet.datavalidation import DataValidation
    path = Path(path)
    if path.exists():
        raise ValueError('Refusing to overwrite an existing review workbook.')
    workbook = Workbook(); sheet = workbook.active; sheet.title = '标注'
    sheet.append(HEADERS)
    labels = {metric: label for metric, _, label, _ in METRICS}
    for row in packet['rows']:
        for metric in CATEGORIES[row['category']][1]:
            sheet.append([row['blind_id'], row['category'], row['text'], metric, labels[metric],
                          None, 'unreviewed', None, None, None, None, None, '否'])
    for cells in sheet:
        for cell in cells:
            if isinstance(cell.value, str):
                cell.data_type = 's'
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for column, width in zip('ABCDEFGHIJKLM', (18,12,62,24,18,16,26,10,24,42,42,24,20)):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = 'F2'; sheet.auto_filter.ref = sheet.dimensions
    for column, choices in (('G', STATES), ('L', ('other_person','past_resolved','negated_claim','ironic_literal')), ('M', ('是','否'))):
        validation = DataValidation(type='list', formula1='"'+','.join(choices)+'"', allow_blank=True)
        validation.errorTitle = '请从列表中选择'; validation.showErrorMessage = True
        sheet.add_data_validation(validation); validation.add(f'{column}2:{column}{sheet.max_row}')
    guide = workbook.create_sheet('填写说明')
    for line in ('每行对应一段原文的一个指标；前五列不要修改。',
                 '先独立标注，不能查看旧答案或另一人的表；完成后填写真实标注者编号，把最后一列改为“是”。',
                 'supported：有当前本人依据，填写 0–100 分；supported_unscored：有依据但程度不明，分数留空。',
                 'explicit_absence：明确没有该症状，分数为 0；不适用于心情好坏、睡眠质量、社交意愿或满意度。',
                 'insufficient_evidence：没有足够依据；disputed：仍有分歧；unreviewed：未完成。这三种分数留空。',
                 '有依据时填写对应时段，并把原文中的依据原样复制到“依据原文”；重复出现时复制更完整、唯一的一段。',
                 '需要排除他人、过去已缓解的描述、被否定说法或反话字面内容时，填写“排除原文”和对应原因。',
                 '空白不等于零分。导入工具检查格式与身份记录，不能代替独立人工复核。'):
        guide.append([line])
    guide.column_dimensions['A'].width = 120
    for cells in guide:
        cells[0].alignment = Alignment(wrap_text=True)
    from .review import review_guide
    scales = workbook.create_sheet('分数方向')
    scales.append(['指标', '0 分含义', '100 分含义'])
    for metric in review_guide()['metrics']:
        scales.append([metric['label'], metric['zero'], metric['hundred']])
    for column in 'ABC':
        scales.column_dimensions[column].width = 32
    workbook.save(path)
    return {'path': str(path), 'cases': len(packet['rows']), 'metric_rows': sheet.max_row-1, 'human_reviews_completed': 0}


def _span(text, quote, reason=None):
    # PSEUDOCODE: resolve only an exact unique copied substring; never guess the intended repeated occurrence.
    if not isinstance(quote, str) or not quote or text.count(quote) != 1:
        raise ValueError('Evidence/exclusion must quote a unique exact part of the original text.')
    start = text.index(quote)
    return {'start': start, 'end': start+len(quote), 'quote': quote, **({'reason': reason} if reason else {})}


def import_review_sheet(packet, path):
    # PSEUDOCODE: preserve the blind inventory -> require explicit independent-human attestation -> validate scores and scope.
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        sheet = workbook['标注']
        entries = list(sheet.values)
    finally:
        workbook.close()
    if not entries or tuple(entries[0]) != HEADERS:
        raise ValueError('Review workbook columns changed.')
    originals = {row['blind_id']: row for row in packet['rows']}
    if len(originals) != len(packet['rows']):
        raise ValueError('Assigned review identities must be unique.')
    expected = {(identity, metric) for identity, row in originals.items() for metric in CATEGORIES[row['category']][1]}
    result, seen, raters = deepcopy(originals), set(), set()
    for row in result.values():
        row['scope_targets'] = {}
    for identity, category, text, metric, label, rater, state, score, period, evidence, excluded, reason, affirmed in entries[1:]:
        key = (identity, metric)
        if key not in expected or key in seen:
            raise ValueError('Review workbook inventory changed or contains duplicate metric rows.')
        seen.add(key)
        original = originals[identity]; row = result[identity]
        metric_label = next(m[2] for m in METRICS if m[0] == metric)
        if text != original['text'] or category != original['category'] or label != metric_label:
            raise ValueError('Review source text, category or metric label changed.')
        if affirmed != '是' or not isinstance(rater, str) or not rater.strip() or state == 'unreviewed':
            raise ValueError('A completed independent human review and a real rater identity must be declared.')
        raters.add(rater.strip())
        row.update(rater_id=rater.strip(), human_reviewed=True)
        row['scores'][metric] = score; row['label_states'][metric] = state
        if state in ('supported', 'supported_unscored', 'explicit_absence') and not evidence:
            raise ValueError('Supported assessments need an exact evidence quote.')
        if reason and not excluded:
            raise ValueError('An exclusion reason needs its source quote.')
        if evidence or excluded:
            if not isinstance(period, str) or not period.strip():
                raise ValueError('Scope annotations need an explicit observation period.')
            row.setdefault('scope_targets', {})[metric] = {'experiencer':'self', 'period':period,
                'evidence': [_span(text, evidence)] if evidence else [],
                'excluded': [_span(text, excluded, reason)] if excluded else []}
    if seen != expected or len(raters) != 1:
        raise ValueError('Complete every assigned metric using one independently identified rater per packet.')
    for row in result.values():
        validate_labels(row); validate_scope_targets(row)
    return {'rows': list(result.values()), 'review_workbook_sha256': file_hash(path),
            'human_review_claim': 'declared_by_named_rater_not_certified_by_importer', 'automatic_adjudication': False}
