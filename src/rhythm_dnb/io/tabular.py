"""Lossless exchange views; display formatting never mutates scientific values."""

import csv
import json
from pathlib import Path
from ..provenance import canonical_json


def read_table(path):
    # PSEUDOCODE: dispatch by explicit format, preserving empty values rather than inventing defaults.
    path = Path(path)
    if path.suffix.lower() == '.json':
        rows = json.loads(path.read_text(encoding='utf-8-sig'))
    elif path.suffix.lower() == '.jsonl':
        rows = [json.loads(line) for line in path.read_text(encoding='utf-8-sig').splitlines() if line.strip()]
    elif path.suffix.lower() == '.csv':
        with path.open(encoding='utf-8-sig', newline='') as stream:
            reader = csv.DictReader(stream)
            header = reader.fieldnames
            if not header or len(set(header)) != len(header) or any(not k.strip() for k in header):
                raise ValueError('Invalid or duplicate CSV headers.')
            rows = list(reader)
            if any(None in row or any(v is None for v in row.values()) for row in rows):
                raise ValueError('CSV row width differs from the declared header.')
    elif path.suffix.lower() == '.xlsx':
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            iterator = workbook.active.iter_rows(values_only=True)
            header = next(iterator, ())
            if not header or len(set(header)) != len(header) or any(not isinstance(k, str) or not k.strip() for k in header):
                raise ValueError('Invalid workbook headers.')
            rows = [dict(zip(header, values)) for values in iterator]
        finally:
            workbook.close()
    else:
        raise ValueError('Supported formats: JSON, JSONL, CSV, XLSX.')
    if not isinstance(rows, list) or any(not isinstance(row, dict) for row in rows):
        raise ValueError('Table must be a list of row objects.')
    return rows


def export_view(rows, path):
    # PSEUDOCODE: create a fresh exchange file without rounding or overwriting another artifact.
    path = Path(path); path.parent.mkdir(parents=True, exist_ok=True)
    if path.suffix.lower() != '.json':
        raise ValueError('The lossless exchange view uses a .json file.')
    with path.open('x', encoding='utf-8') as stream:
        stream.write(canonical_json(list(rows)) + '\n')
    return path
