"""Blind spreadsheet roundtrip and rejection of missing, guessed or altered human evidence."""

from pathlib import Path
import tempfile
import unittest
from openpyxl import load_workbook
from rhythm_dnb.text.review_sheet import export_review_sheet, import_review_sheet


class ReviewSheetTests(unittest.TestCase):
    def test_blank_form_then_explicit_completed_review(self):
        packet = {'rows': [{'blind_id': 'software-only', 'category': 'stress', 'text': '今天我压力很大，同事说他还好。',
            'scores': {'stress_intensity': None}, 'label_states': {'stress_intensity': 'unreviewed'}, 'scope_targets': {}}]}
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder)/'review.xlsx'
            export_review_sheet(packet, path)
            with self.assertRaisesRegex(ValueError, 'overwrite'):
                export_review_sheet(packet, path)
            with self.assertRaisesRegex(ValueError, 'completed independent human review'):
                import_review_sheet(packet, path)
            workbook = load_workbook(path); sheet = workbook['标注']
            for column, value in {'F':'software-test-rater', 'G':'supported', 'H':80., 'I':'今天',
                                  'J':'我压力很大', 'K':'同事说他还好', 'L':'other_person', 'M':'是'}.items():
                sheet[column+'2'] = value
            workbook.save(path)
            result = import_review_sheet(packet, path)
            self.assertEqual(result['rows'][0]['scores']['stress_intensity'], 80)
            self.assertEqual(result['rows'][0]['scope_targets']['stress_intensity']['evidence'][0]['start'], 2)
            self.assertIn('not_certified', result['human_review_claim'])
            self.assertIsNone(packet['rows'][0]['scores']['stress_intensity'])
            sheet['C2'] = '修改了原文'; workbook.save(path)
            with self.assertRaisesRegex(ValueError, 'source text'):
                import_review_sheet(packet, path)

    def test_repeated_quotes_cannot_be_silently_assigned(self):
        packet = {'rows': [{'blind_id':'software-only', 'category':'stress', 'text':'他压力很大，我压力很大。',
            'scores': {'stress_intensity':None}, 'label_states':{'stress_intensity':'unreviewed'}}]}
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder)/'review.xlsx'; export_review_sheet(packet, path)
            workbook = load_workbook(path); sheet = workbook['标注']
            for column, value in {'F':'test-rater','G':'supported_unscored','I':'当前','J':'压力很大','M':'是'}.items():
                sheet[column+'2'] = value
            workbook.save(path)
            with self.assertRaisesRegex(ValueError, 'unique exact'):
                import_review_sheet(packet, path)
            sheet['J2'] = '我压力很大'; sheet['H2'] = 50; workbook.save(path)
            with self.assertRaises(ValueError):
                import_review_sheet(packet, path)
