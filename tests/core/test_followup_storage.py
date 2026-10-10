"""Simulated follow-up must be recoverable, isolated and fully exported."""

from contextlib import closing, redirect_stdout
import io
from pathlib import Path
import sqlite3
import tempfile
import unittest
from zipfile import ZipFile

from rhythm_dnb.provenance import file_hash
from tools.complete_research_followup import prepare, apply, workbook


class FollowupStorageTests(unittest.TestCase):
    def fixture(self, root):
        database = root / 'data.sqlite'
        with closing(sqlite3.connect(database)) as db, db:
            db.execute('CREATE TABLE observations(record_id TEXT PRIMARY KEY,participant_id TEXT,source_dataset TEXT,'
                       'observed_at TEXT,split TEXT,sleep_start_hour REAL,sleep_end_hour REAL,sleep_duration_h REAL,'
                       'exercise_hour REAL,meal_interval_cv REAL,exercise_minutes REAL,resting_hr_bpm REAL,screen_time_min REAL)')
            db.execute('CREATE TABLE raw_inputs(record_id TEXT PRIMARY KEY,payload TEXT)')
            db.execute('CREATE TABLE observation_provenance(record_id TEXT PRIMARY KEY,payload TEXT)')
            db.execute('CREATE TABLE research_annotation_runs(run_id TEXT PRIMARY KEY)')
            db.execute('INSERT INTO research_annotation_runs VALUES (\'initial\')')
            db.execute('CREATE TABLE research_outcome_annotations(run_id TEXT,record_id TEXT,candidate_state INTEGER,'
                       'reference_event_in_followup INTEGER,reference_onset_at TEXT,is_simulated_cohort INTEGER,status TEXT,reference_scope TEXT)')
            for i in range(30):
                person = f'p{i:02d}'; rid = person+'-one'
                db.execute('INSERT INTO observations VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)',
                           (rid,person,'source','2020-01-01T12:00:00+00:00','train',23,7,8,17,.1,35,60,120))
                for table in ('raw_inputs','observation_provenance'):
                    db.execute(f'INSERT INTO {table} VALUES (?,?)', (rid,'{}'))
                db.execute('INSERT INTO research_outcome_annotations VALUES (?,?,?,?,?,?,?,?)',
                           ('initial',rid,None,None,None,0,'baseline_only','not_provided'))
        output = root / 'proposal'
        with redirect_stdout(io.StringIO()):
            prepare(database, output)
        return database, output

    def test_preserves_old_rows_exact_backup_parent_roles_and_export(self):
        from openpyxl import Workbook, load_workbook
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder); database, output = self.fixture(root)
            before = file_hash(database)
            receipt = apply(database, output)
            self.assertEqual(file_hash(output/'rhythm.before.sqlite'), before)
            self.assertTrue(receipt['all_original_tables_unchanged'])
            self.assertEqual(receipt['added_records'], 2520)
            with closing(sqlite3.connect(database)) as db:
                self.assertEqual(db.execute('SELECT count(*) FROM research_all_record_status').fetchone()[0], 2550)
                self.assertEqual(db.execute('SELECT count(*) FROM research_followup_anchors').fetchone()[0], 30)
                self.assertEqual(db.execute('SELECT sum(reference_event_in_followup) FROM research_outcome_annotations').fetchone()[0], None)
                self.assertEqual(db.execute('SELECT count(*) FROM research_followup_days WHERE features_json IS NULL').fetchone()[0], 0)
            path = root/'quantified_data.xlsx'
            original = Workbook(); original.active.title = '量化总表'; original.active.append(['保持原文',123]); original.save(path); original.close()
            with ZipFile(path) as z:
                xml = z.read('xl/worksheets/sheet1.xml')
            report = workbook(database, path, output)
            with ZipFile(path) as z:
                self.assertEqual(z.read('xl/worksheets/sheet1.xml'), xml)
            check = load_workbook(path, read_only=True)
            self.assertEqual(check['量化总表']['A1'].value, '保持原文')
            self.assertEqual(check['模拟随访记录'].max_row, 2521)
            self.assertEqual(check['补齐概览'].max_row, 31)
            check.close()
            self.assertTrue(report['original_worksheet_xml_unchanged'])
            with self.assertRaisesRegex(ValueError, 'Database changed'):
                apply(database, output)

    def test_changed_original_blocks_application_without_backup_or_writes(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder); database, output = self.fixture(root)
            with closing(sqlite3.connect(database)) as db, db:
                db.execute('UPDATE observations SET exercise_minutes=99 WHERE participant_id=\'p00\'')
            with self.assertRaisesRegex(ValueError, 'Database changed'):
                apply(database, output)
            self.assertFalse((output/'rhythm.before.sqlite').exists())
            with closing(sqlite3.connect(database)) as db:
                self.assertEqual(db.execute("SELECT count(*) FROM sqlite_master WHERE name LIKE 'research_followup_%'").fetchone()[0], 0)

    def test_altered_labels_are_rejected_before_writing(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder); database, output = self.fixture(root)
            path = output/'sequences.json'
            path.write_text(path.read_text(encoding='utf-8')+' ',encoding='utf-8')
            with self.assertRaisesRegex(ValueError, 'Changed proposal artifact'):
                apply(database, output)
            self.assertFalse((output/'rhythm.before.sqlite').exists())


if __name__ == '__main__':
    unittest.main()
