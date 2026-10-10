"""Append auditable candidate labels and verified warnings to the existing dataset."""

import argparse
from collections import Counter
from contextlib import closing
from datetime import datetime, timezone
import json
import math
from pathlib import Path
import re
import shutil
import sqlite3
import xml.etree.ElementTree as ET
from xml.sax.saxutils import escape
from zipfile import ZIP_DEFLATED, ZipFile

from rhythm_dnb.provenance import canonical_json, file_hash, fingerprint
from rhythm_dnb.research.dataset_annotations import POLICY, annotate_records
from rhythm_dnb.research.experiment_data import read_snapshot, save_json
from rhythm_dnb.io.sources.lineage import CONSTRUCTED_METHODS


METHODS = ('reference_robust_dnb', 'reference_robust_deviation', 'history_robust_dnb', 'history_robust_control')
STATUS = {'baseline_only': '基线记录，不判事件', 'insufficient_baseline': '基线不足',
          'followup_window_incomplete': '后续窗口不足', 'insufficient_window': '当日或窗口数据不足',
          'no_candidate_worsening': '未达到候选事件条件', 'candidate_confirmed': '规则候选事件',
          'persistence_not_confirmed': '变化持续天数不足', 'existing_scenario_preserved': '沿用场景参考答案'}


def read_json(path):
    # PSEUDOCODE: read a previously saved UTF-8 JSON artifact.
    return json.loads(Path(path).read_text(encoding='utf-8'))


def verified_warnings(directory, source):
    # PSEUDOCODE: verify the saved replay inventory, then join every method by exact record/person identity.
    directory = Path(directory)
    manifest = read_json(directory / 'manifest.json')
    if fingerprint(manifest['files']) != manifest['id']:
        raise ValueError('Warning replay manifest is invalid.')
    for name, digest in manifest['files'].items():
        if Path(name).name != name or file_hash(directory / name) != digest:
            raise ValueError('Warning artifact changed: ' + name)
    verification = read_json(directory / 'verification.json')
    if verification.get('event_metrics_match') is not True or verification.get('rows_replayed') != 3600:
        raise ValueError('Warning replay was not verified.')
    known = {r['record_id']: r for r in source}
    result, seen = [], set()
    for fold in range(1, 4):
        for method in METHODS:
            name = f'fold-{fold}-{method}-predictions.json'
            saved = read_json(directory / name)
            if saved['method'] != method:
                raise ValueError('Warning method mismatch.')
            for row in saved['rows']:
                original = known.get(row['record_id'])
                if original is None or original['participant_id'] != row['participant_id'] or original['source_dataset'] != 'SIMULATED-RHYTHM':
                    raise ValueError('Warning has no matching scenario record.')
                key = (row['record_id'], method)
                if key in seen or row['warning'] not in (None, 0, 1):
                    raise ValueError('Repeated or invalid warning decision.')
                seen.add(key)
                if row['score'] is not None and (type(row['score']) not in (float, int) or not math.isfinite(row['score'])):
                    raise ValueError('Warning score must be finite or missing.')
                if datetime.fromisoformat(row['issued_at']) < datetime.fromisoformat(original['observed_at']):
                    raise ValueError('Warning precedes its input record.')
                result.append({**row, 'method': method, 'fold': fold, 'bundle_id': saved['bundle_id'],
                               'artifact': name, 'artifact_sha256': manifest['files'][name]})
    if len(result) != verification['rows_replayed']:
        raise ValueError('Incomplete warning replay.')
    return result, manifest['id']


def prepare(database, replay, output):
    # PSEUDOCODE: freeze the rule and source identity before computing any new candidate labels.
    output = Path(output)
    output.mkdir(parents=True, exist_ok=True)
    source, source_hash = read_snapshot(database)
    plan = {'created_at': datetime.now(timezone.utc).isoformat(), 'database': str(Path(database).resolve()),
            'source_sha256': source_hash, 'policy': POLICY, 'records': len(source),
            'annotation_engine': 'deterministic_rules_designed_by_assistant_not_per_record_LLM_inference',
            'label_selection_uses_warning_accuracy': False, 'original_labels_preserved': True}
    plan['id'] = fingerprint(plan)
    if (output / 'plan.json').exists():
        if {p.name for p in output.iterdir()} != {'plan.json'}:
            raise FileExistsError('A proposal already exists; do not overwrite it.')
        plan = read_json(output / 'plan.json')
        if (plan['source_sha256'] != source_hash or plan['policy'] != POLICY or
                plan['id'] != fingerprint({k: v for k, v in plan.items() if k != 'id'})):
            raise ValueError('Cannot resume a changed source or annotation policy.')
    else:
        if any(output.iterdir()):
            raise FileExistsError('Use an empty proposal directory.')
        save_json(output / 'plan.json', plan)
    annotations, streams = annotate_records(source)
    originals = {r['record_id']: r for r in source}
    for annotation in annotations:
        original = originals[annotation['record_id']]
        provenance = json.loads(original['provenance']).get('provenance', {})
        annotation['contains_constructed_values'] = int(annotation['is_simulated_cohort'] or any(
            isinstance(v, dict) and v.get('method') in CONSTRUCTED_METHODS for v in provenance.values()))
    warnings, replay_id = verified_warnings(replay, source)
    save_json(output / 'annotations.json', annotations)
    save_json(output / 'sequence-evidence.json', streams)
    save_json(output / 'warnings.json', warnings)
    states = Counter(str(r['candidate_state']) for r in annotations)
    candidate_people = {(s['source_dataset'], s['participant_id']) for s in streams if s['events']}
    summary = {'source_sha256': source_hash, 'annotation_run_id': plan['id'], 'records': len(annotations),
               'candidate_states': dict(states), 'annotation_statuses': dict(Counter(r['status'] for r in annotations)),
               'candidate_event_count': sum(len(s['events']) for s in streams),
               'candidate_source_person_sequences': len(candidate_people),
               'candidate_participant_ids': len({p for _, p in candidate_people}),
               'warnings': dict(Counter(r['method'] for r in warnings)),
               'alarms': dict(Counter(r['method'] for r in warnings if r['warning'] == 1)),
               'saved_warning_replay_id': replay_id, 'new_clinical_gold_labels': 0,
               'original_measurement_or_reference_label_changes': 0,
               'scopes': dict(Counter(r['reference_scope'] for r in annotations))}
    save_json(output / 'summary.json', summary)
    files = {p.name: file_hash(p) for p in output.glob('*.json')}
    save_json(output / 'manifest.json', {'files': files, 'id': fingerprint(files)})
    if file_hash(database) != source_hash:
        raise ValueError('Database changed during annotation; do not apply this proposal.')
    return summary


def _original_table_hashes(connection, tables):
    # PSEUDOCODE: stream every original row and schema so adding annotations cannot silently alter any old table.
    import hashlib
    result = {}
    for name in tables:
        quoted = '"' + name.replace('"', '""') + '"'
        digest = hashlib.sha256()
        schema = connection.execute('SELECT sql FROM sqlite_master WHERE type=\'table\' AND name=?', (name,)).fetchone()[0]
        digest.update(schema.encode('utf-8'))
        count = 0
        for row in connection.execute(f'SELECT * FROM {quoted} ORDER BY rowid'):
            digest.update(canonical_json(list(row)).encode('utf-8')); digest.update(b'\n'); count += 1
        result[name] = {'rows': count, 'sha256': digest.hexdigest()}
    return result


def apply(database, output):
    # PSEUDOCODE: preserve a byte-verified source snapshot, append only new tables, verify all old tables before commit.
    database, output = Path(database).resolve(), Path(output)
    manifest = read_json(output / 'manifest.json')
    if fingerprint(manifest['files']) != manifest['id']:
        raise ValueError('Invalid annotation manifest.')
    for name, digest in manifest['files'].items():
        if Path(name).name != name or file_hash(output / name) != digest:
            raise ValueError('Annotation proposal changed: ' + name)
    plan, annotations, streams, warnings = [read_json(output / p) for p in
        ('plan.json', 'annotations.json', 'sequence-evidence.json', 'warnings.json')]
    if fingerprint({k: v for k, v in plan.items() if k != 'id'}) != plan['id'] or file_hash(database) != plan['source_sha256']:
        raise ValueError('Plan or source changed; no in-place overwrite allowed.')
    backup = output / 'rhythm.before.sqlite'
    if backup.exists():
        raise FileExistsError('Inspect existing backup before retrying.')
    # Obtain the write lock before backup so a concurrent dataset edit cannot race the snapshot.
    with closing(sqlite3.connect(database)) as db, db:
        if db.execute('PRAGMA journal_mode').fetchone()[0] != 'delete':
            raise ValueError('Byte-identical source backup requires a checkpointed DELETE-mode source.')
        db.execute('PRAGMA foreign_keys=ON'); db.execute('BEGIN IMMEDIATE')
        if file_hash(database) != plan['source_sha256']:
            raise ValueError('Source changed before write lock.')
        shutil.copy2(database, backup)
        if file_hash(backup) != plan['source_sha256']:
            raise ValueError('Backup hash mismatch.')
        tables = [r[0] for r in db.execute("SELECT name FROM sqlite_master WHERE type='table' ORDER BY name")]
        before = _original_table_hashes(db, tables)
        db.execute('CREATE TABLE research_annotation_runs (run_id TEXT PRIMARY KEY, created_at TEXT NOT NULL, '
                   'source_sha256 TEXT NOT NULL, policy_json TEXT NOT NULL, sequence_evidence_json TEXT NOT NULL)')
        db.execute('CREATE TABLE research_outcome_annotations (run_id TEXT NOT NULL REFERENCES research_annotation_runs(run_id), '
                   'record_id TEXT NOT NULL REFERENCES observations(record_id), is_simulated_cohort INTEGER NOT NULL, '
                   'contains_constructed_values INTEGER NOT NULL, reference_scope TEXT NOT NULL, '
                   'reference_event_in_followup INTEGER, reference_onset_at TEXT, candidate_state INTEGER, '
                   'candidate_onset_date TEXT, candidate_confirmed_date TEXT, status TEXT NOT NULL, evidence_json TEXT NOT NULL, '
                   'PRIMARY KEY(run_id,record_id), CHECK(candidate_state IS NULL OR candidate_state IN (0,1)))')
        db.execute('CREATE TABLE research_warning_annotations (run_id TEXT NOT NULL REFERENCES research_annotation_runs(run_id), '
                   'record_id TEXT NOT NULL REFERENCES observations(record_id), method TEXT NOT NULL, issued_at TEXT NOT NULL, '
                   'warning INTEGER, score REAL, status TEXT NOT NULL, payload_json TEXT NOT NULL, '
                   'PRIMARY KEY(run_id,record_id,method))')
        db.execute('INSERT INTO research_annotation_runs VALUES (?,?,?,?,?)',
                   (plan['id'], plan['created_at'], plan['source_sha256'], canonical_json(plan['policy']), canonical_json(streams)))
        for row in annotations:
            values = [row[k] for k in ('record_id','is_simulated_cohort','contains_constructed_values','reference_scope',
                      'reference_event_in_followup','reference_onset_at','candidate_state','candidate_onset_date','candidate_confirmed_date','status')]
            db.execute('INSERT INTO research_outcome_annotations VALUES (?,?,?,?,?,?,?,?,?,?,?,?)',
                       [plan['id'], *values, canonical_json(row['evidence'])])
        for row in warnings:
            db.execute('INSERT INTO research_warning_annotations VALUES (?,?,?,?,?,?,?,?)',
                       [plan['id'], *[row[k] for k in ('record_id','method','issued_at','warning','score','status')], canonical_json(row)])
        db.execute('CREATE VIEW research_labeled_observations AS SELECT o.*, a.is_simulated_cohort, a.contains_constructed_values, '
                   'a.reference_scope, a.reference_event_in_followup, a.candidate_state, a.candidate_onset_date, '
                   'a.candidate_confirmed_date, a.status AS annotation_status, a.run_id AS annotation_run_id '
                   'FROM observations o LEFT JOIN research_outcome_annotations a USING(record_id)')
        after = _original_table_hashes(db, tables)
        if before != after or db.execute('PRAGMA foreign_key_check').fetchone():
            raise ValueError('Original data changed or new references invalid; transaction rolled back.')
        if db.execute('PRAGMA quick_check').fetchone()[0] != 'ok':
            raise ValueError('Database integrity check failed.')
    receipt = {'annotation_run_id': plan['id'], 'database': str(database), 'backup': str(backup),
               'database_before_sha256': plan['source_sha256'], 'database_after_sha256': file_hash(database),
               'all_original_tables_unchanged': True, 'original_tables': before,
               'annotation_rows_added': len(annotations), 'saved_warning_rows_added': len(warnings)}
    save_json(output / 'database-verification.json', receipt)
    return receipt


def update_workbook(workbook, output):
    # PSEUDOCODE: append readable annotation columns, preserve every old cell and verify the updated file before replacement.
    from openpyxl.utils import get_column_letter
    from openpyxl import load_workbook
    workbook, output = Path(workbook), Path(output)
    if not (output / 'database-verification.json').exists():
        raise ValueError('Apply and verify the database before updating its workbook view.')
    receipt = read_json(output / 'database-verification.json')
    if file_hash(receipt['database']) != receipt['database_after_sha256']:
        raise ValueError('Database changed since annotation; do not publish a stale workbook view.')
    manifest = read_json(output / 'manifest.json')
    if fingerprint(manifest['files']) != manifest['id']:
        raise ValueError('Invalid annotation manifest.')
    for name, digest in manifest['files'].items():
        if Path(name).name != name or file_hash(output / name) != digest:
            raise ValueError('Annotation proposal changed: ' + name)
    annotations = {r['record_id']: r for r in read_json(output / 'annotations.json')}
    warnings = {(r['record_id'], r['method']): r for r in read_json(output / 'warnings.json')}
    columns = [('数据组成', 'annotation_data_origin'), ('已有参考标签类型', 'annotation_reference_scope'),
               ('随访内是否发生（已有答案）', 'reference_event_in_followup'),
               ('候选变化状态（规则初标）', 'candidate_state'), ('候选变化起始日期', 'candidate_onset_date'),
               ('候选变化确认日期', 'candidate_confirmed_date'), ('初标说明', 'annotation_status'),
               ('单独DNB预警（0/1）', 'pure_dnb_warning'), ('单独DNB预警计算时间', 'pure_dnb_issued_at'),
               ('单独DNB分数', 'pure_dnb_score'), ('DNB组合预警（0/1）', 'joint_warning'),
               ('DNB组合预警计算时间', 'joint_issued_at'), ('DNB组合分数', 'joint_score'),
               ('单独DNB预警状态', 'pure_dnb_status'), ('DNB组合预警状态', 'joint_status')]
    scopes = {'scenario': '场景设定', 'prevalent_sleep_phase_diagnosis': '入组时睡眠相位延迟诊断', 'not_provided': '无既有结局答案'}
    before_hash = file_hash(workbook)
    backup = output / 'quantified_data.before.xlsx'
    if backup.exists():
        raise FileExistsError('Workbook already has an annotation backup.')
    shutil.copy2(workbook, backup)
    with ZipFile(workbook) as z:
        entries = {name: z.read(name) for name in z.namelist()}
    sheet = entries['xl/worksheets/sheet1.xml'].decode('utf-8')
    if 'annotation_data_origin' in sheet:
        raise ValueError('Workbook is already annotated.')
    # Read identities through openpyxl; never infer them from row positions or abbreviated identifiers.
    book = load_workbook(workbook, read_only=True)
    width = book['量化总表'].max_column
    identities = [r[0] for r in book['量化总表'].iter_rows(min_row=2, values_only=True)]
    book.close()
    if len(identities) != len(set(identities)) or set(identities) != set(annotations):
        raise ValueError('Workbook and database record identities differ.')
    all_values = [[label + '\n' + key for label, key in columns]]
    for rid in identities:
        row = annotations[rid]
        pure, joint = (warnings.get((rid, m), {}) for m in ('reference_robust_dnb', 'history_robust_dnb'))
        origin = '模拟队列' if row['is_simulated_cohort'] else '含补全指标的实验记录' if row['contains_constructed_values'] else '来源记录'
        all_values.append([origin, scopes[row['reference_scope']], row['reference_event_in_followup'],
            row['candidate_state'], row['candidate_onset_date'], row['candidate_confirmed_date'], STATUS[row['status']],
            pure.get('warning'), pure.get('issued_at'), pure.get('score'), joint.get('warning'), joint.get('issued_at'),
            joint.get('score'), pure.get('status', '未在已完成实验中评估'), joint.get('status', '未在已完成实验中评估')])

    def append_cells(match):
        number = int(re.search(r'\br="(\d+)"', match.group(0)).group(1))
        cells = []
        for offset, value in enumerate(all_values[number - 1], width + 1):
            ref = f'{get_column_letter(offset)}{number}'
            if value is None:
                continue
            if isinstance(value, str):
                cells.append(f'<c r="{ref}" t="inlineStr"><is><t>{escape(value)}</t></is></c>')
            else:
                cells.append(f'<c r="{ref}" t="n"><v>{value}</v></c>')
        return match.group(0)[:-6] + ''.join(cells) + '</row>'

    updated = re.sub(r'<row\b[^>]*>.*?</row>', append_cells, sheet, flags=re.DOTALL)
    end = f'{get_column_letter(width + len(columns))}{len(identities) + 1}'
    updated = re.sub(r'(<dimension ref=")[^"]+("/>)', lambda m: m[1] + 'A1:' + end + m[2], updated)
    updated = re.sub(r'(<autoFilter ref=")[^"]+("[^>]*>)', lambda m: m[1] + 'A1:' + end + m[2], updated)
    widths = ''.join(f'<col min="{i}" max="{i}" width="25" customWidth="1"/>' for i in range(width + 1, width + len(columns) + 1))
    updated = updated.replace('</cols>', widths + '</cols>')
    entries['xl/worksheets/sheet1.xml'] = updated.encode('utf-8')
    # Existing summary stays in the same workbook; add a compact explanation of the newly visible fields.
    note = ('候选状态是按固定规则生成的初标，0仅表示当日未满足相对变化条件；起始日期精确到天。'
            '预警0表示已计算但未发出信号，空白表示未评估。单独DNB与DNB组合分别显示。'
            '候选初标不能用作独立测试金标准；已有参考答案和全部原始字段保留。')
    if 'xl/worksheets/sheet2.xml' in entries:
        summary = entries['xl/worksheets/sheet2.xml'].decode('utf-8')
        summary = summary.replace('</sheetData>', f'<row r="6"><c r="A6" t="inlineStr"><is><t>{escape(note)}</t></is></c></row></sheetData>')
        entries['xl/worksheets/sheet2.xml'] = summary.encode('utf-8')
    # Summarize people without turning an unassessed person into a negative case.
    from collections import defaultdict
    people = defaultdict(list)
    for row in annotations.values():
        people[row['participant_id']].append(row)
    events = defaultdict(list)
    for stream in read_json(output / 'sequence-evidence.json'):
        events[stream['participant_id']].extend(stream['events'])
    alerts = defaultdict(list)
    for row in warnings.values():
        alerts[(row['participant_id'], row['method'])].append(row)
    overview = [['受试者编号', '记录数', '数据组成', '随访内是否发生（已有答案）', '发生时间（已有答案）',
                 '规则初标可判断天数', '候选事件段数', '首段候选起始日期', '首段候选确认日期',
                 '单独DNB已计算次数', '单独DNB预警次数', '单独DNB首次预警时间',
                 'DNB组合已计算次数', 'DNB组合预警次数', 'DNB组合首次预警时间', '标签说明']]
    for person, items in sorted(people.items()):
        simulated = all(r['is_simulated_cohort'] for r in items)
        reference = {r['reference_event_in_followup'] for r in items if r['reference_event_in_followup'] is not None}
        if len(reference) > 1:
            raise ValueError('Inconsistent reference outcomes for one person.')
        candidate = sorted(events[person], key=lambda e: e['onset_date'])
        assessed = sum(r['candidate_state'] is not None for r in items)
        row = [person, len(items), '模拟队列' if simulated else '含补全指标的实验记录',
               next(iter(reference), None), min((r['reference_onset_at'] for r in items if r['reference_onset_at']), default=None),
               assessed, len(candidate) if assessed else None,
               candidate[0]['onset_date'] if candidate else None, candidate[0]['confirmed_date'] if candidate else None]
        for method in ('reference_robust_dnb', 'history_robust_dnb'):
            predictions = alerts[(person, method)]
            fired = [r for r in predictions if r['warning'] == 1]
            row.extend([len(predictions), len(fired) if predictions else None,
                        min((r['issued_at'] for r in fired), default=None)])
        row.append('场景设定' if reference else '单次稳定参考，不作未来事件阴性随访' if simulated else
                   '候选标注；原入组睡眠诊断不作新发事件' if any(r['reference_scope']=='prevalent_sleep_phase_diagnosis' for r in items) else
                   '候选标注；没有独立未来结局答案')
        overview.append(row)
    main_ns = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
    doc_rel = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
    rel_ns = 'http://schemas.openxmlformats.org/package/2006/relationships'
    content_ns = 'http://schemas.openxmlformats.org/package/2006/content-types'
    book_xml = ET.fromstring(entries['xl/workbook.xml'])
    sheets = book_xml.find(f'{{{main_ns}}}sheets')
    sheet_id = max(int(s.attrib['sheetId']) for s in sheets) + 1
    file_name, relation = f'xl/worksheets/sheet{sheet_id}.xml', 'rIdAnnotationOverview'
    if file_name in entries or any(s.attrib['name'] == '事件与预警' for s in sheets):
        raise ValueError('Annotation overview already exists.')
    ET.SubElement(sheets, f'{{{main_ns}}}sheet', {'name':'事件与预警','sheetId':str(sheet_id),f'{{{doc_rel}}}id':relation})
    relationships = ET.fromstring(entries['xl/_rels/workbook.xml.rels'])
    ET.SubElement(relationships,f'{{{rel_ns}}}Relationship',{'Type':doc_rel+'/worksheet','Target':'/'+file_name,'Id':relation})
    types = ET.fromstring(entries['[Content_Types].xml'])
    ET.SubElement(types,f'{{{content_ns}}}Override',{'PartName':'/'+file_name,
        'ContentType':'application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml'})
    entries['xl/workbook.xml'] = ET.tostring(book_xml,encoding='utf-8')
    entries['xl/_rels/workbook.xml.rels'] = ET.tostring(relationships,encoding='utf-8')
    entries['[Content_Types].xml'] = ET.tostring(types,encoding='utf-8')
    overview_rows = []
    for i, values in enumerate(overview,1):
        cells = []
        for j, value in enumerate(values,1):
            if value is None:
                continue
            ref = f'{get_column_letter(j)}{i}'
            cells.append(f'<c r="{ref}" t="inlineStr"><is><t>{escape(value)}</t></is></c>' if isinstance(value,str)
                         else f'<c r="{ref}" t="n"><v>{value}</v></c>')
        overview_rows.append(f'<row r="{i}">'+''.join(cells)+'</row>')
    overview_end = f'P{len(overview)}'
    entries[file_name] = (f'<worksheet xmlns="{main_ns}"><dimension ref="A1:{overview_end}"/>'
        '<sheetViews><sheetView workbookViewId="0"><pane xSplit="1" ySplit="1" topLeftCell="B2" state="frozen"/></sheetView></sheetViews>'
        '<cols><col min="1" max="1" width="44" customWidth="1"/><col min="2" max="16" width="30" customWidth="1"/></cols>'
        '<sheetData>'+''.join(overview_rows)+f'</sheetData><autoFilter ref="A1:{overview_end}"/></worksheet>').encode('utf-8')
    pending = output / 'quantified_data.pending.xlsx'
    with ZipFile(pending, 'w', ZIP_DEFLATED) as z:
        for name, payload in entries.items():
            z.writestr(name, payload)
    # Strip appended cells from the staged XML and compare the original row XML byte for byte.
    for old, new in zip(re.findall(r'<row\b[^>]*>.*?</row>', sheet, re.DOTALL),
                        re.findall(r'<row\b[^>]*>.*?</row>', updated, re.DOTALL)):
        original_prefix = old[:-6]
        if not new.startswith(original_prefix):
            raise ValueError('An original workbook cell changed.')
    check = load_workbook(pending, read_only=True)
    assert check['量化总表'].max_row == len(identities) + 1
    assert check['量化总表'].max_column == width + len(columns)
    check.close()
    if file_hash(workbook) != before_hash:
        raise ValueError('Workbook changed concurrently; pending file retained.')
    pending.replace(workbook)
    receipt = {'workbook': str(workbook), 'records': len(identities), 'columns_added': columns,
               'original_cells_unchanged': True, 'before_sha256': before_hash, 'after_sha256': file_hash(workbook)}
    save_json(output / 'workbook-verification.json', receipt)
    return receipt


def main():
    # PSEUDOCODE: require explicit preparation, application or workbook projection; never silently overwrite a proposal.
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('stage', choices=('prepare', 'apply', 'workbook'))
    parser.add_argument('--database', type=Path)
    parser.add_argument('--replay', type=Path)
    parser.add_argument('--workbook', type=Path)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    if args.stage == 'prepare':
        if not args.database or not args.replay:
            parser.error('prepare requires --database and --replay')
        result = prepare(args.database, args.replay, args.output)
    elif args.stage == 'apply':
        if not args.database:
            parser.error('apply requires --database')
        result = apply(args.database, args.output)
    else:
        if not args.workbook:
            parser.error('workbook requires --workbook')
        result = update_workbook(args.workbook, args.output)
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
